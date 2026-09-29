"""Nexora PDF Tools - 30 self-hosted PDF utilities.

Every handler takes (files, opts) and returns (data, download_name, mime).
files: list of (filename, bytes). opts: dict of str -> str from the form.
Raise ToolError with a user-safe message for expected failures.
"""

import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile


class ToolError(Exception):
    pass


# ----------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------

def _one(files, what="a PDF file"):
    if not files:
        raise ToolError("Please upload " + what + " first.")
    return files[0]


def _pdf_reader(data, password=None):
    from pypdf import PdfReader
    try:
        r = PdfReader(io.BytesIO(data))
        if r.is_encrypted:
            ok = r.decrypt(password or "")
            if not ok:
                raise ToolError("That PDF is locked with a password. "
                                "Unlock it first with the correct password.")
        return r
    except ToolError:
        raise
    except Exception:
        raise ToolError("Could not read that PDF. It may be damaged - try Repair PDF first.")


def _writer_bytes(w):
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _parse_ranges(spec, n):
    """'1-3,5' -> list of 0-based page indexes. Raises ToolError."""
    out = []
    spec = (spec or "").strip()
    if not spec:
        raise ToolError("Enter at least one page or range, like 1-3 or 2,4,6.")
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a > b:
                a, b = b, a
            out.extend(range(a - 1, b))
        elif part.isdigit():
            out.append(int(part) - 1)
        else:
            raise ToolError("Could not understand '" + part + "'. Use pages like 1-3,5.")
    out = [i for i in out if 0 <= i < n]
    if not out:
        raise ToolError("None of those pages exist - the document has "
                        + str(n) + " page(s).")
    return out


def _zip_named(parts):
    """parts: list of (arcname, bytes) -> zip bytes."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts:
            z.writestr(name, data)
    return buf.getvalue()


def _soffice():
    """Locate a LibreOffice binary (portable install or PATH)."""
    for cand in (os.environ.get("NEXORA_SOFFICE", ""),
                 "/opt/render/project/lo/program/soffice",
                 shutil.which("soffice") or "",
                 shutil.which("libreoffice") or ""):
        if cand and os.path.exists(cand):
            return cand
    return None


def _lo_convert(data, in_name, timeout=180):
    """Convert an office document to PDF via LibreOffice. Returns pdf bytes."""
    soffice = _soffice()
    if not soffice:
        raise ToolError("office-converter-unavailable")
    work = tempfile.mkdtemp(prefix="nexora_lo_")
    try:
        src = os.path.join(work, in_name)
        with open(src, "wb") as f:
            f.write(data)
        subprocess.run(
            [soffice, "--headless", "--norestore", "--convert-to",
             "pdf", "--outdir", work, src],
            check=True, capture_output=True, timeout=timeout)
        out = os.path.splitext(src)[0] + ".pdf"
        with open(out, "rb") as f:
            return f.read()
    except subprocess.TimeoutExpired:
        raise ToolError("That file took too long to convert. Try a smaller file.")
    except Exception:
        raise ToolError("Could not convert that file. Check that it opens "
                        "correctly on your computer and try again.")
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _rl_canvas_overlay(w, h, draw):
    """Build a one-page PDF overlay of size (w,h) via reportlab. draw(canvas)."""
    from reportlab.pdfgen import canvas as rl_canvas
    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=(w, h))
    draw(c)
    c.save()
    buf.seek(0)
    from pypdf import PdfReader
    return PdfReader(buf).pages[0]


def _apply_overlay(reader, page_filter, overlay_for):
    """Merge overlay onto pages; overlay_for(i, page) -> overlay page or None."""
    from pypdf import PdfWriter
    w = PdfWriter()
    for i, page in enumerate(reader.pages):
        if page_filter(i):
            ov = overlay_for(i, page)
            if ov is not None:
                page.merge_page(ov)
        w.add_page(page)
    return _writer_bytes(w)


# ----------------------------------------------------------------------
# 1-9: merge / split / compress / conversions
# ----------------------------------------------------------------------

def h_merge(files, opts):
    from pypdf import PdfWriter
    if len(files) < 2:
        raise ToolError("Upload at least two PDF files to merge.")
    w = PdfWriter()
    for name, data in files:
        r = _pdf_reader(data)
        for p in r.pages:
            w.add_page(p)
    return _writer_bytes(w), "merged.pdf", "application/pdf"


def h_split(files, opts):
    name, data = _one(files)
    r = _pdf_reader(data)
    n = len(r.pages)
    spec = (opts.get("ranges") or "").strip()
    from pypdf import PdfWriter
    if not spec:
        parts = []
        for i in range(n):
            w = PdfWriter()
            w.add_page(r.pages[i])
            parts.append(("page-%d.pdf" % (i + 1), _writer_bytes(w)))
        return _zip_named(parts), "split-pages.zip", "application/zip"
    groups = [g.strip() for g in spec.split(";") if g.strip()]
    if len(groups) == 1:
        idxs = _parse_ranges(groups[0], n)
        w = PdfWriter()
        for i in idxs:
            w.add_page(r.pages[i])
        return _writer_bytes(w), "split.pdf", "application/pdf"
    parts = []
    for gi, g in enumerate(groups):
        idxs = _parse_ranges(g, n)
        w = PdfWriter()
        for i in idxs:
            w.add_page(r.pages[i])
        parts.append(("part-%d.pdf" % (gi + 1), _writer_bytes(w)))
    return _zip_named(parts), "split-parts.zip", "application/zip"


def h_compress(files, opts):
    name, data = _one(files)
    level = (opts.get("level") or "balanced").lower()
    try:
        import pikepdf
        pdf = pikepdf.open(io.BytesIO(data))
        buf = io.BytesIO()
        pdf.save(buf, compress_streams=True, object_stream_mode=pikepdf.ObjectStreamMode.generate,
                 recompress_flate=True)
        out = buf.getvalue()
    except ToolError:
        raise
    except Exception:
        raise ToolError("Could not compress that PDF. It may be damaged - try Repair PDF first.")
    if level in ("strong", "extreme"):
        try:
            import fitz
            doc = fitz.open(stream=out, filetype="pdf")
            dpi = 110 if level == "strong" else 85
            new = fitz.open()
            for page in doc:
                pix = page.get_pixmap(dpi=dpi)
                p = new.new_page(width=page.rect.width, height=page.rect.height)
                p.insert_image(p.rect, pixmap=pix)
            buf2 = io.BytesIO()
            new.save(buf2, garbage=4, deflate=True)
            smaller = buf2.getvalue()
            if len(smaller) < len(out):
                out = smaller
        except Exception:
            pass
    return out, "compressed.pdf", "application/pdf"


def h_pdf_to_word(files, opts):
    name, data = _one(files)
    work = tempfile.mkdtemp(prefix="nexora_p2w_")
    try:
        from pdf2docx import Converter
        srcf = os.path.join(work, "in.pdf")
        dst = os.path.join(work, "out.docx")
        with open(srcf, "wb") as f:
            f.write(data)
        cv = Converter(srcf)
        cv.convert(dst)
        cv.close()
        with open(dst, "rb") as f:
            out = f.read()
        if not out:
            raise ValueError("empty output")
        return out, "converted.docx", \
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    except Exception:
        raise ToolError("Could not convert that PDF to Word. Scanned PDFs "
                        "(images of pages) need OCR instead.")
    finally:
        shutil.rmtree(work, ignore_errors=True)


def h_word_to_pdf(files, opts):
    name, data = _one(files)
    low = name.lower()
    try:
        return _lo_convert(data, os.path.basename(name)), "converted.pdf", "application/pdf"
    except ToolError as e:
        if "office-converter-unavailable" not in str(e):
            raise
    # fallback: render docx content (text, headings, tables) with reportlab
    if not low.endswith(".docx"):
        raise ToolError("Please upload a .docx Word file.")
    try:
        import docx
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas as rl_canvas
        d = docx.Document(io.BytesIO(data))
        buf = io.BytesIO()
        c = rl_canvas.Canvas(buf, pagesize=A4)
        W, H = A4
        y = H - 20 * mm
        def line(text, size=11, bold=False):
            nonlocal y
            if y < 20 * mm:
                c.showPage()
                y = H - 20 * mm
            c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
            for chunk in _wrap(text, 95):
                c.drawString(20 * mm, y, chunk)
                y -= size * 1.5
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                y -= 6
                continue
            style = (p.style.name or "").lower()
            line(t, 15 if "heading" in style else 11, "heading" in style or "title" in style)
        for tbl in d.tables:
            for row in tbl.rows:
                line(" | ".join(cell.text.strip() for cell in row.cells), 10)
            y -= 4
        c.save()
        return buf.getvalue(), "converted.pdf", "application/pdf"
    except Exception:
        raise ToolError("Could not convert that Word file. Check that it opens "
                        "correctly and try again.")


def _wrap(text, width):
    words, cur = text.split(), ""
    for w in words:
        if len(cur) + len(w) + 1 > width:
            yield cur
            cur = w
        else:
            cur = (cur + " " + w).strip()
    if cur:
        yield cur


def h_pdf_to_ppt(files, opts):
    name, data = _one(files)
    try:
        import fitz
        from pptx import Presentation
        from pptx.util import Inches
        doc = fitz.open(stream=data, filetype="pdf")
        prs = Presentation()
        blank = prs.slide_layouts[6]
        for page in doc:
            pix = page.get_pixmap(dpi=120)
            img = io.BytesIO(pix.tobytes("png"))
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(img, 0, 0, width=prs.slide_width,
                                     height=prs.slide_height)
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue(), "converted.pptx", \
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    except Exception:
        raise ToolError("Could not convert that PDF to PowerPoint.")


def h_ppt_to_pdf(files, opts):
    name, data = _one(files)
    try:
        return _lo_convert(data, os.path.basename(name)), "converted.pdf", "application/pdf"
    except ToolError as e:
        if "office-converter-unavailable" not in str(e):
            raise
    try:
        from pptx import Presentation
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas as rl_canvas
        prs = Presentation(io.BytesIO(data))
        buf = io.BytesIO()
        W, H = landscape(A4)
        c = rl_canvas.Canvas(buf, pagesize=(W, H))
        for slide in prs.slides:
            y = H - 20 * mm
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = "".join(run.text for run in para.runs).strip()
                        if t:
                            c.setFont("Helvetica", 13)
                            for chunk in _wrap(t, 110):
                                c.drawString(15 * mm, y, chunk)
                                y -= 8 * mm
                                if y < 15 * mm:
                                    c.showPage()
                                    y = H - 20 * mm
            c.showPage()
        c.save()
        return buf.getvalue(), "converted.pdf", "application/pdf"
    except Exception:
        raise ToolError("Could not convert that PowerPoint file.")


def h_pdf_to_excel(files, opts):
    name, data = _one(files)
    try:
        import pdfplumber
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        found = False
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for pi, page in enumerate(pdf.pages):
                tables = page.extract_tables() or []
                if tables:
                    found = True
                    ws = wb.create_sheet("page-%d" % (pi + 1))
                    for tbl in tables:
                        for row in tbl:
                            ws.append([("" if c is None else str(c)) for c in row])
                        ws.append([])
        if not found:
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            ws = wb.create_sheet("text")
            ws.append(["Page", "Line"])
            for pi, page in enumerate(doc):
                for ln in page.get_text().splitlines():
                    if ln.strip():
                        ws.append([pi + 1, ln])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "converted.xlsx", \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    except Exception:
        raise ToolError("Could not convert that PDF to Excel.")


def h_excel_to_pdf(files, opts):
    name, data = _one(files)
    try:
        return _lo_convert(data, os.path.basename(name)), "converted.pdf", "application/pdf"
    except ToolError as e:
        if "office-converter-unavailable" not in str(e):
            raise
    try:
        from openpyxl import load_workbook
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas as rl_canvas
        wb = load_workbook(io.BytesIO(data), data_only=True)
        buf = io.BytesIO()
        W, H = A4
        c = rl_canvas.Canvas(buf, pagesize=A4)
        y = H - 18 * mm
        c.setFont("Helvetica", 9)
        for ws in wb.worksheets:
            c.setFont("Helvetica-Bold", 13)
            c.drawString(15 * mm, y, ws.title[:60])
            y -= 10 * mm
            c.setFont("Helvetica", 9)
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if not any(vals):
                    continue
                txt = " | ".join(vals)[:160]
                c.drawString(15 * mm, y, txt)
                y -= 5.5 * mm
                if y < 15 * mm:
                    c.showPage()
                    y = H - 18 * mm
                    c.setFont("Helvetica", 9)
            c.showPage()
            y = H - 18 * mm
        c.save()
        return buf.getvalue(), "converted.pdf", "application/pdf"
    except Exception:
        raise ToolError("Could not convert that Excel file.")


# ----------------------------------------------------------------------
# 10-16: images, edit, sign, watermark, rotate, html
# ----------------------------------------------------------------------

def h_pdf_to_jpg(files, opts):
    name, data = _one(files)
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        parts = []
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=150)
            parts.append(("page-%d.jpg" % (i + 1), pix.tobytes("jpeg", jpg_quality=85)))
        if len(parts) == 1:
            return parts[0][1], "page-1.jpg", "image/jpeg"
        return _zip_named(parts), "pages-jpg.zip", "application/zip"
    except Exception:
        raise ToolError("Could not convert that PDF to images.")


def _images_to_pdf(files, scan=False):
    from PIL import Image, ImageOps
    pages = []
    for name, data in files:
        try:
            img = Image.open(io.BytesIO(data))
            img = ImageOps.exif_transpose(img)
            if scan:
                img = ImageOps.grayscale(img)
                img = ImageOps.autocontrast(img, cutoff=2)
            if img.mode != "RGB":
                img = img.convert("RGB")
            pages.append(img)
        except Exception:
            raise ToolError("One of those files is not a readable image.")
    if not pages:
        raise ToolError("Upload at least one image.")
    buf = io.BytesIO()
    pages[0].save(buf, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    return buf.getvalue()


def h_jpg_to_pdf(files, opts):
    return _images_to_pdf(files, scan=False), "images.pdf", "application/pdf"


def h_scan_to_pdf(files, opts):
    return _images_to_pdf(files, scan=True), "scan.pdf", "application/pdf"


def h_edit(files, opts):
    name, data = _one(files)
    text = (opts.get("text") or "").strip()
    if not text:
        raise ToolError("Type the text to add.")
    page_no = int(opts.get("page") or "1")
    x = min(max(float(opts.get("x") or "10"), 0), 95) / 100.0
    y = min(max(float(opts.get("y") or "10"), 0), 95) / 100.0
    size = min(max(float(opts.get("size") or "18"), 6), 120)
    r = _pdf_reader(data)
    n = len(r.pages)
    if not (1 <= page_no <= n):
        raise ToolError("That document has " + str(n) + " page(s) - pick a page in range.")
    target = page_no - 1
    def ov(i, page):
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        def draw(c):
            c.setFont("Helvetica-Bold", size)
            c.setFillColorRGB(0.1, 0.1, 0.35)
            for j, ln in enumerate(text.split("\n")):
                c.drawString(w * x, h * (1 - y) - j * size * 1.4, ln)
        return _rl_canvas_overlay(w, h, draw)
    return _apply_overlay(r, lambda i: i == target, ov), "edited.pdf", "application/pdf"


def h_sign(files, opts):
    if not files:
        raise ToolError("Upload a PDF first.")
    name, data = files[0]
    sig_img = None
    if len(files) > 1:
        sig_img = files[1][1]
    text = (opts.get("text") or "").strip()
    if not sig_img and not text:
        raise ToolError("Type your name for the signature, or attach a signature image.")
    page_no = int(opts.get("page") or "0")  # 0 = last page
    x = min(max(float(opts.get("x") or "60"), 0), 90) / 100.0
    y = min(max(float(opts.get("y") or "85"), 0), 98) / 100.0
    r = _pdf_reader(data)
    n = len(r.pages)
    target = (n - 1) if page_no == 0 else page_no - 1
    if not (0 <= target < n):
        raise ToolError("That document has " + str(n) + " page(s) - pick a page in range.")
    def ov(i, page):
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        def draw(c):
            if sig_img:
                from reportlab.lib.utils import ImageReader
                img = ImageReader(io.BytesIO(sig_img))
                iw, ih = img.getSize()
                dw = w * 0.28
                dh = dw * ih / max(iw, 1)
                c.drawImage(img, w * x, h * (1 - y) - dh, dw, dh,
                            mask="auto", preserveAspectRatio=True)
            else:
                c.setFont("Helvetica-Oblique", 26)
                c.setFillColorRGB(0.05, 0.05, 0.4)
                c.drawString(w * x, h * (1 - y) - 26, text)
        return _rl_canvas_overlay(w, h, draw)
    return _apply_overlay(r, lambda i: i == target, ov), "signed.pdf", "application/pdf"


def h_watermark(files, opts):
    name, data = _one(files)
    text = (opts.get("text") or "CONFIDENTIAL").strip()[:60]
    r = _pdf_reader(data)
    def ov(i, page):
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        def draw(c):
            c.saveState()
            c.translate(w / 2, h / 2)
            c.rotate(45)
            c.setFont("Helvetica-Bold", 60)
            c.setFillColorRGB(0.6, 0.1, 0.3, alpha=0.18)
            c.drawCentredString(0, 0, text)
            c.restoreState()
        return _rl_canvas_overlay(w, h, draw)
    return _apply_overlay(r, lambda i: True, ov), "watermarked.pdf", "application/pdf"


def h_rotate(files, opts):
    name, data = _one(files)
    angle = int(opts.get("angle") or "90")
    if angle not in (90, 180, 270):
        raise ToolError("Rotation must be 90, 180 or 270 degrees.")
    r = _pdf_reader(data)
    from pypdf import PdfWriter
    w = PdfWriter()
    for p in r.pages:
        p.rotate(angle)
        w.add_page(p)
    return _writer_bytes(w), "rotated.pdf", "application/pdf"


def h_html_to_pdf(files, opts):
    name, data = _one(files)
    try:
        html = data.decode("utf-8", errors="replace")
    except Exception:
        raise ToolError("Could not read that file as HTML text.")
    try:
        from xhtml2pdf import pisa
        buf = io.BytesIO()
        res = pisa.CreatePDF(io.StringIO(html), dest=buf)
        if res.err:
            raise ValueError("parse error")
        return buf.getvalue(), "webpage.pdf", "application/pdf"
    except Exception:
        raise ToolError("Could not convert that HTML. Complex pages with remote "
                        "scripts may not convert - try a simpler page.")


# ----------------------------------------------------------------------
# 17-22: unlock, protect, organize, repair, page numbers
# ----------------------------------------------------------------------

def h_unlock(files, opts):
    name, data = _one(files)
    from pypdf import PdfReader, PdfWriter
    try:
        r = PdfReader(io.BytesIO(data))
        if r.is_encrypted:
            ok = r.decrypt(opts.get("password") or "")
            if not ok:
                raise ToolError("Wrong password - the PDF stayed locked.")
        w = PdfWriter()
        for p in r.pages:
            w.add_page(p)
        return _writer_bytes(w), "unlocked.pdf", "application/pdf"
    except ToolError:
        raise
    except Exception:
        raise ToolError("Could not unlock that PDF.")


def h_protect(files, opts):
    name, data = _one(files)
    pw = (opts.get("password") or "").strip()
    if len(pw) < 4:
        raise ToolError("Choose a password of at least 4 characters.")
    r = _pdf_reader(data)
    from pypdf import PdfWriter
    w = PdfWriter()
    for p in r.pages:
        w.add_page(p)
    w.encrypt(pw)
    return _writer_bytes(w), "protected.pdf", "application/pdf"


def h_organize(files, opts):
    name, data = _one(files)
    r = _pdf_reader(data)
    n = len(r.pages)
    order = (opts.get("order") or "").strip()
    if not order:
        raise ToolError("Enter the new page order, like 3,1,2 or 1,4-6. "
                        "Pages you leave out are removed.")
    idxs = _parse_ranges(order, n)
    from pypdf import PdfWriter
    w = PdfWriter()
    for i in idxs:
        w.add_page(r.pages[i])
    return _writer_bytes(w), "organized.pdf", "application/pdf"


def h_repair(files, opts):
    name, data = _one(files)
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        buf = io.BytesIO()
        doc.save(buf, garbage=4, deflate=True)
        return buf.getvalue(), "repaired.pdf", "application/pdf"
    except Exception:
        pass
    try:
        import pikepdf
        pdf = pikepdf.open(io.BytesIO(data))
        buf = io.BytesIO()
        pdf.save(buf)
        return buf.getvalue(), "repaired.pdf", "application/pdf"
    except Exception:
        raise ToolError("That PDF is too damaged to repair.")


def h_page_numbers(files, opts):
    name, data = _one(files)
    pos = (opts.get("position") or "bottom-center").lower()
    r = _pdf_reader(data)
    n = len(r.pages)
    def ov(i, page):
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        def draw(c):
            c.setFont("Helvetica", 11)
            c.setFillColorRGB(0.25, 0.25, 0.3)
            label = str(i + 1) + " / " + str(n)
            x = w / 2 if "center" in pos else (w - 40 if "right" in pos else 40)
            y = 24 if "bottom" in pos else h - 24
            c.drawCentredString(x, y, label) if "center" in pos else c.drawString(x - 20, y, label)
        return _rl_canvas_overlay(w, h, draw)
    return _apply_overlay(r, lambda i: True, ov), "numbered.pdf", "application/pdf"


# ----------------------------------------------------------------------
# 23-30: ocr, compare, redact, crop, markdown, txt, images, grayscale
# ----------------------------------------------------------------------

def _tesseract_ready():
    if shutil.which("tesseract"):
        return True
    td = os.environ.get("TESSDATA_PREFIX", "")
    return bool(td and os.path.exists(os.path.join(td, "eng.traineddata")))


def _ocr_rapidocr(data):
    """Searchable PDF via RapidOCR (ONNX, no system deps). Returns pdf bytes."""
    import cv2
    import fitz
    import numpy as np
    from rapidocr import RapidOCR
    engine = RapidOCR()
    src = fitz.open(stream=data, filetype="pdf")
    out = fitz.open()
    for page in src:
        pix = page.get_pixmap(dpi=200)
        newp = out.new_page(width=page.rect.width, height=page.rect.height)
        newp.insert_image(newp.rect, pixmap=pix)
        img = cv2.imdecode(np.frombuffer(pix.tobytes("png"), np.uint8),
                           cv2.IMREAD_COLOR)
        result = engine(img)
        if result is None or result.txts is None:
            continue
        sx = page.rect.width / pix.width
        sy = page.rect.height / pix.height
        for box, text, conf in zip(result.boxes, result.txts, result.scores):
            if conf < 0.5 or not text.strip():
                continue
            box = [[float(px), float(py)] for px, py in box]
            x0 = min(p[0] for p in box) * sx
            y0 = min(p[1] for p in box) * sy
            y1 = max(p[1] for p in box) * sy
            size = max(4.0, (y1 - y0) * 0.9)
            newp.insert_text(fitz.Point(x0, y1 - size * 0.15), text,
                             fontsize=size, render_mode=3)
    buf = io.BytesIO()
    out.save(buf, garbage=3, deflate=True)
    return buf.getvalue()


def h_ocr(files, opts):
    name, data = _one(files)
    if not _tesseract_ready():
        try:
            return _ocr_rapidocr(data), "ocr-searchable.pdf", "application/pdf"
        except ImportError:
            raise ToolError("ocr-unavailable")
        except Exception:
            raise ToolError("OCR failed on that document. Try a clearer scan.")
    try:
        import fitz
        src = fitz.open(stream=data, filetype="pdf")
        out = fitz.open()
        for page in src:
            pix = page.get_pixmap(dpi=200)
            pdfbytes = pix.pdfocr_tobytes(language="eng")
            ocr_doc = fitz.open("pdf", pdfbytes)
            out.insert_pdf(ocr_doc)
        buf = io.BytesIO()
        out.save(buf, garbage=3, deflate=True)
        return buf.getvalue(), "ocr-searchable.pdf", "application/pdf"
    except Exception:
        raise ToolError("OCR failed on that document. Try a clearer scan.")


def h_compare(files, opts):
    if len(files) < 2:
        raise ToolError("Upload TWO PDFs to compare (older first, newer second).")
    import difflib
    import fitz
    texts = []
    for name, data in files[:2]:
        try:
            doc = fitz.open(stream=data, filetype="pdf")
            texts.append("\n".join(p.get_text() for p in doc).splitlines())
        except Exception:
            raise ToolError("Could not read one of those PDFs.")
    diff = difflib.HtmlDiff(wrapcolumn=90)
    html = diff.make_file(texts[0], texts[1],
                          fromdesc=files[0][0], todesc=files[1][0], context=True)
    return html.encode("utf-8"), "comparison.html", "text/html"


def h_redact(files, opts):
    name, data = _one(files)
    words = [w.strip() for w in (opts.get("words") or "").split(",") if w.strip()]
    if not words:
        raise ToolError("Enter the word(s) or phrase(s) to remove, comma-separated.")
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        hits = 0
        for page in doc:
            for w in words:
                areas = page.search_for(w)
                for a in areas:
                    page.add_redact_annot(a, fill=(0, 0, 0))
                    hits += 1
            page.apply_redactions()
        if not hits:
            raise ToolError("None of those words were found in the document.")
        buf = io.BytesIO()
        doc.save(buf, garbage=4, deflate=True)
        return buf.getvalue(), "redacted.pdf", "application/pdf"
    except ToolError:
        raise
    except Exception:
        raise ToolError("Could not redact that PDF.")


def h_crop(files, opts):
    name, data = _one(files)
    try:
        m = max(float(opts.get("margin") or "10"), 0)
    except ValueError:
        raise ToolError("Margin must be a number of millimetres.")
    mm = 72 / 25.4
    r = _pdf_reader(data)
    from pypdf import PdfWriter
    w = PdfWriter()
    off = m * mm
    for p in r.pages:
        box = p.mediabox
        if box.width - 2 * off < 50 or box.height - 2 * off < 50:
            raise ToolError("That margin is too large for this page size.")
        box.left += off
        box.right -= off
        box.bottom += off
        box.top -= off
        w.add_page(p)
    return _writer_bytes(w), "cropped.pdf", "application/pdf"


def h_pdf_to_markdown(files, opts):
    name, data = _one(files)
    try:
        import pymupdf4llm
        md = pymupdf4llm.to_markdown(io.BytesIO(data))
        if md.strip():
            return md.encode("utf-8"), "document.md", "text/markdown"
    except Exception:
        pass
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        parts = []
        for i, page in enumerate(doc):
            parts.append("\n## Page %d\n" % (i + 1))
            parts.append(page.get_text())
        return "\n".join(parts).encode("utf-8"), "document.md", "text/markdown"
    except Exception:
        raise ToolError("Could not convert that PDF to Markdown.")


def h_pdf_to_txt(files, opts):
    name, data = _one(files)
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        txt = "\n\n".join(p.get_text() for p in doc)
        if not txt.strip():
            raise ToolError("That PDF has no selectable text - try OCR PDF for scans.")
        return txt.encode("utf-8"), "document.txt", "text/plain"
    except ToolError:
        raise
    except Exception:
        raise ToolError("Could not extract text from that PDF.")


def h_extract_images(files, opts):
    name, data = _one(files)
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        parts = []
        seen = set()
        for pi, page in enumerate(doc):
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                ext = doc.extract_image(xref)
                parts.append(("image-p%d-%d.%s" % (pi + 1, xref, ext["ext"]),
                              ext["image"]))
        if not parts:
            raise ToolError("No embedded images found in that PDF.")
        if len(parts) == 1:
            nm, blob = parts[0]
            return blob, nm, "image/" + nm.rsplit(".", 1)[-1].replace("jpg", "jpeg")
        return _zip_named(parts), "extracted-images.zip", "application/zip"
    except ToolError:
        raise
    except Exception:
        raise ToolError("Could not extract images from that PDF.")


def h_grayscale(files, opts):
    name, data = _one(files)
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        new = fitz.open()
        for page in doc:
            pix = page.get_pixmap(dpi=150, colorspace=fitz.csGRAY)
            p = new.new_page(width=page.rect.width, height=page.rect.height)
            p.insert_image(p.rect, pixmap=pix)
        buf = io.BytesIO()
        new.save(buf, garbage=4, deflate=True)
        return buf.getvalue(), "grayscale.pdf", "application/pdf"
    except Exception:
        raise ToolError("Could not convert that PDF to grayscale.")


# ----------------------------------------------------------------------
# registry
# ----------------------------------------------------------------------

PDF_MIME = "application/pdf"

TOOLS = {
    "merge": dict(title="Merge PDF", icon="merge",
                  desc="Combine PDFs in the order you want into one file.",
                  accept=".pdf", multiple=True, options=[], handler=h_merge),
    "split": dict(title="Split PDF", icon="split",
                  desc="Extract a page range, split into parts, or one file per page.",
                  accept=".pdf", multiple=False,
                  options=[dict(key="ranges", label="Pages (optional)",
                                ph="e.g. 1-3,5 - separate parts with ; - empty = every page")],
                  handler=h_split),
    "compress": dict(title="Compress PDF", icon="compress",
                     desc="Shrink file size while keeping the best quality possible.",
                     accept=".pdf", multiple=False,
                     options=[dict(key="level", label="Compression", kind="select",
                                   choices=[("balanced", "Balanced"), ("strong", "Strong"),
                                            ("extreme", "Extreme (smallest)")])],
                     handler=h_compress),
    "pdf-to-word": dict(title="PDF to Word", icon="word",
                        desc="Turn a PDF into an editable DOCX document.",
                        accept=".pdf", multiple=False, options=[], handler=h_pdf_to_word),
    "word-to-pdf": dict(title="Word to PDF", icon="word",
                        desc="Make DOC and DOCX files easy to read by converting them to PDF.",
                        accept=".doc,.docx", multiple=False, options=[], handler=h_word_to_pdf),
    "pdf-to-ppt": dict(title="PDF to PowerPoint", icon="ppt",
                       desc="Turn each PDF page into a PowerPoint slide.",
                       accept=".pdf", multiple=False, options=[], handler=h_pdf_to_ppt),
    "ppt-to-pdf": dict(title="PowerPoint to PDF", icon="ppt",
                       desc="Make PPT and PPTX slideshows easy to view as PDF.",
                       accept=".ppt,.pptx", multiple=False, options=[], handler=h_ppt_to_pdf),
    "pdf-to-excel": dict(title="PDF to Excel", icon="excel",
                         desc="Pull tables and text from a PDF into a spreadsheet.",
                         accept=".pdf", multiple=False, options=[], handler=h_pdf_to_excel),
    "excel-to-pdf": dict(title="Excel to PDF", icon="excel",
                         desc="Convert XLS and XLSX spreadsheets to a clean PDF.",
                         accept=".xls,.xlsx,.csv", multiple=False, options=[], handler=h_excel_to_pdf),
    "pdf-to-jpg": dict(title="PDF to JPG", icon="image",
                       desc="Turn every PDF page into a JPG image.",
                       accept=".pdf", multiple=False, options=[], handler=h_pdf_to_jpg),
    "jpg-to-pdf": dict(title="JPG to PDF", icon="image",
                       desc="Convert JPG / PNG images into a single PDF.",
                       accept=".jpg,.jpeg,.png,.webp", multiple=True, options=[], handler=h_jpg_to_pdf),
    "edit": dict(title="Edit PDF", icon="pen",
                 desc="Add text anywhere on a page - notes, labels, corrections.",
                 accept=".pdf", multiple=False,
                 options=[dict(key="text", label="Text to add", ph="Type the text"),
                          dict(key="page", label="Page number", ph="1"),
                          dict(key="x", label="X position % (0=left, 90=right)", ph="10"),
                          dict(key="y", label="Y position % (0=top, 90=bottom)", ph="10"),
                          dict(key="size", label="Font size", ph="18")],
                 handler=h_edit),
    "sign": dict(title="Sign PDF", icon="sign",
                 desc="Stamp your signature - type your name or attach a signature image.",
                 accept=".pdf", multiple=True,
                 options=[dict(key="text", label="Your name (typed signature)", ph="A. Dennis"),
                          dict(key="page", label="Page (0 = last page)", ph="0"),
                          dict(key="x", label="X position % (0=left)", ph="60"),
                          dict(key="y", label="Y position % (0=top)", ph="85")],
                 extra_file="Signature image (optional, PNG/JPG)", handler=h_sign),
    "watermark": dict(title="Watermark PDF", icon="stamp",
                      desc="Stamp diagonal text across every page.",
                      accept=".pdf", multiple=False,
                      options=[dict(key="text", label="Watermark text", ph="CONFIDENTIAL")],
                      handler=h_watermark),
    "rotate": dict(title="Rotate PDF", icon="rotate",
                   desc="Turn every page 90, 180 or 270 degrees.",
                   accept=".pdf", multiple=False,
                   options=[dict(key="angle", label="Rotation", kind="select",
                                 choices=[("90", "90° clockwise"), ("180", "180°"),
                                          ("270", "270° clockwise")])],
                   handler=h_rotate),
    "html-to-pdf": dict(title="HTML to PDF", icon="code",
                        desc="Convert a saved webpage (.html file) into a PDF.",
                        accept=".html,.htm,.txt", multiple=False, options=[], handler=h_html_to_pdf),
    "unlock": dict(title="Unlock PDF", icon="unlock",
                   desc="Remove password protection from a PDF you own.",
                   accept=".pdf", multiple=False,
                   options=[dict(key="password", label="Current password (if any)", ph="leave empty if none")],
                   handler=h_unlock),
    "protect": dict(title="Protect PDF", icon="lock",
                    desc="Lock a PDF with a password so only you can open it.",
                    accept=".pdf", multiple=False,
                    options=[dict(key="password", label="New password", ph="at least 4 characters")],
                    handler=h_protect),
    "organize": dict(title="Organize PDF", icon="organize",
                     desc="Reorder or delete pages - set the exact page order.",
                     accept=".pdf", multiple=False,
                     options=[dict(key="order", label="New page order", ph="e.g. 3,1,2,5-6 - left-out pages are deleted")],
                     handler=h_organize),
    "repair": dict(title="Repair PDF", icon="wrench",
                   desc="Recover data from a damaged or corrupt PDF.",
                   accept=".pdf", multiple=False, options=[], handler=h_repair),
    "page-numbers": dict(title="Page Numbers", icon="numbers",
                         desc="Add page numbers to every page, where you want them.",
                         accept=".pdf", multiple=False,
                         options=[dict(key="position", label="Position", kind="select",
                                       choices=[("bottom-center", "Bottom center"),
                                                ("bottom-right", "Bottom right"),
                                                ("bottom-left", "Bottom left"),
                                                ("top-center", "Top center"),
                                                ("top-right", "Top right"),
                                                ("top-left", "Top left")])],
                         handler=h_page_numbers),
    "scan-to-pdf": dict(title="Scan to PDF", icon="scan",
                        desc="Turn photos of paper pages into a clean, enhanced PDF scan.",
                        accept=".jpg,.jpeg,.png,.webp", multiple=True, options=[], handler=h_scan_to_pdf),
    "ocr": dict(title="OCR PDF", icon="ocr",
                desc="Make a scanned PDF searchable and selectable.",
                accept=".pdf", multiple=False, options=[], handler=h_ocr),
    "compare": dict(title="Compare PDF", icon="compare",
                    desc="Upload two versions and see every change highlighted.",
                    accept=".pdf", multiple=True, options=[],
                    extra_file="Second PDF (newer version)", handler=h_compare),
    "redact": dict(title="Redact PDF", icon="redact",
                   desc="Permanently black out sensitive words or phrases.",
                   accept=".pdf", multiple=False,
                   options=[dict(key="words", label="Words to remove", ph="comma, separated, phrases")],
                   handler=h_redact),
    "crop": dict(title="Crop PDF", icon="crop",
                 desc="Trim an even margin off every page.",
                 accept=".pdf", multiple=False,
                 options=[dict(key="margin", label="Margin to trim (mm)", ph="10")],
                 handler=h_crop),
    "pdf-to-markdown": dict(title="PDF to Markdown", icon="markdown",
                            desc="Convert a PDF into clean Markdown text.",
                            accept=".pdf", multiple=False, options=[], handler=h_pdf_to_markdown),
    "pdf-to-txt": dict(title="PDF to Text", icon="txt",
                       desc="Extract all selectable text from a PDF.",
                       accept=".pdf", multiple=False, options=[], handler=h_pdf_to_txt),
    "extract-images": dict(title="Extract Images", icon="images",
                           desc="Pull every embedded image out of a PDF.",
                           accept=".pdf", multiple=False, options=[], handler=h_extract_images),
    "grayscale": dict(title="Grayscale PDF", icon="gray",
                      desc="Convert a colour PDF to black and white for printing.",
                      accept=".pdf", multiple=False, options=[], handler=h_grayscale),
}

ORDER = ["merge", "split", "compress", "pdf-to-word", "word-to-pdf",
         "pdf-to-ppt", "ppt-to-pdf", "pdf-to-excel", "excel-to-pdf",
         "pdf-to-jpg", "jpg-to-pdf", "edit", "sign", "watermark", "rotate",
         "html-to-pdf", "unlock", "protect", "organize", "repair",
         "page-numbers", "scan-to-pdf", "ocr", "compare", "redact", "crop",
         "pdf-to-markdown", "pdf-to-txt", "extract-images", "grayscale"]

assert len(ORDER) == 30 and set(ORDER) == set(TOOLS)


def run_tool(slug, files, opts):
    spec = TOOLS.get(slug)
    if not spec:
        raise ToolError("Unknown tool.")
    if not files or not any(d for _, d in files):
        raise ToolError("Please upload a file first.")
    return spec["handler"](files, opts)
