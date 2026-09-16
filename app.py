import base64
import io
import json
import logging
import re

import pandas as pd
import streamlit as st
from google import genai
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

APP_NAME = "NEXORA"
logger = logging.getLogger("nexora")
MODEL = "gemini-3.5-flash-lite"
MAX_PDF_SIZE = 50 * 1024 * 1024
MAX_IMAGE_SIZE = 20 * 1024 * 1024
MAX_AI_FILE_SIZE = 20 * 1024 * 1024

st.set_page_config(page_title="Nexora — AI Document Workspace", page_icon="✦", layout="wide", initial_sidebar_state="collapsed")

# ============================================================
# NEXORA VISUAL SYSTEM — native Streamlit UI, no visible HTML
# ============================================================
st.markdown(r"""
<style>
:root{--blue:#2563eb;--indigo:#4f46e5;--violet:#7c3aed;--cyan:#0891b2;--green:#16a34a;--orange:#f97316;--pink:#ec4899;--ink:#252637;--muted:#687083}
.stApp{background:radial-gradient(circle at 4% 0%,rgba(59,130,246,.23),transparent 24%),radial-gradient(circle at 96% 1%,rgba(236,72,153,.18),transparent 24%),radial-gradient(circle at 82% 42%,rgba(124,58,237,.13),transparent 27%),radial-gradient(circle at 15% 88%,rgba(6,182,212,.14),transparent 28%),linear-gradient(135deg,#f1f8ff 0%,#fbf4ff 46%,#eefbff 100%);min-height:100vh}
.block-container{max-width:1500px!important;padding-top:.15rem!important;padding-bottom:.6rem!important;padding-left:1.1rem!important;padding-right:1.1rem!important}
[data-testid="stHeader"]{background:transparent!important;height:0!important}[data-testid="stToolbar"]{visibility:hidden;height:0}
[data-testid="stVerticalBlock"]{gap:.22rem}.element-container{margin-bottom:.03rem!important}
h1,h2,h3{color:var(--ink)!important}h1{font-weight:900!important;letter-spacing:-2.4px!important}h2{font-weight:850!important;letter-spacing:-1.2px!important}h3{font-weight:800!important}
p,li{color:#51596a}hr{border-color:rgba(148,163,184,.22)!important;margin:.22rem 0 .45rem!important}
[data-testid="stVerticalBlockBorderWrapper"]{background:rgba(255,255,255,.92)!important;border:1px solid rgba(148,163,184,.24)!important;border-radius:18px!important;box-shadow:0 8px 28px rgba(37,52,90,.055),inset 0 1px 0 rgba(255,255,255,.96);transition:transform .18s ease,box-shadow .18s ease,border-color .18s ease}
[data-testid="stVerticalBlockBorderWrapper"]:hover{border-color:rgba(79,70,229,.28)!important;box-shadow:0 14px 36px rgba(79,70,229,.12),inset 0 1px 0 rgba(255,255,255,.98)}
[data-testid="stFileUploader"]{background:linear-gradient(135deg,rgba(239,246,255,.98),rgba(250,245,255,.98))!important;border:2px dashed rgba(79,70,229,.36)!important;border-radius:18px!important;padding:.5rem!important;box-shadow:0 0 0 5px rgba(79,70,229,.035),0 14px 35px rgba(79,70,229,.075);transition:all .2s ease}
[data-testid="stFileUploader"]:hover{border-color:rgba(79,70,229,.72)!important;transform:translateY(-2px);box-shadow:0 0 0 6px rgba(79,70,229,.045),0 18px 42px rgba(79,70,229,.12)}
.stButton>button,.stDownloadButton>button{min-height:40px!important;border-radius:11px!important;font-weight:750!important;border:1px solid rgba(148,163,184,.28)!important;background:rgba(255,255,255,.93)!important;box-shadow:0 4px 14px rgba(37,52,90,.045);transition:all .16s ease}
.stButton>button:hover,.stDownloadButton>button:hover{border-color:rgba(79,70,229,.42)!important;box-shadow:0 8px 22px rgba(79,70,229,.10)}
.stButton>button{position:relative!important;z-index:20!important;pointer-events:auto!important;line-height:1.2!important}
.workspace-choice{min-height:180px!important}.workspace-choice button{min-height:150px!important;font-size:1.08rem!important;border-radius:20px!important}
.stButton>button[kind="primary"],.stButton>button[kind="primary"] p,.stButton>button[kind="primary"] span,.stButton>button[kind="primary"] div{color:#fff!important}.stButton>button[kind="primary"]{color:#fff!important;border:none!important;background:linear-gradient(100deg,#2563eb,#4f46e5 48%,#8b5cf6)!important;box-shadow:0 9px 25px rgba(79,70,229,.24)}
.stButton>button[kind="primary"]:hover{background:linear-gradient(100deg,#1d4ed8,#4338ca 48%,#7c3aed)!important;box-shadow:0 12px 30px rgba(79,70,229,.30)}
div[role="radiogroup"]{gap:.3rem!important;padding:.18rem!important;background:rgba(255,255,255,.72);border:1px solid rgba(148,163,184,.2);border-radius:14px;box-shadow:0 5px 20px rgba(37,52,90,.035)}
div[role="radiogroup"] label{border-radius:10px!important;padding:.3rem .78rem!important;font-weight:700!important}
[data-testid="stChatMessage"]{border-radius:14px!important;margin-bottom:.42rem!important}[data-testid="stChatInput"]{border-radius:14px!important;box-shadow:0 8px 28px rgba(79,70,229,.12)!important}
[data-testid="stDataFrame"]{border-radius:12px!important;overflow:hidden!important;box-shadow:0 6px 20px rgba(37,52,90,.055)}[data-testid="stAlert"]{border-radius:12px!important}.payment-cta{border:1px solid rgba(37,99,235,.22);background:linear-gradient(135deg,rgba(239,246,255,.98),rgba(250,245,255,.98));border-radius:16px;padding:.55rem .7rem;box-shadow:0 10px 28px rgba(79,70,229,.10)}.privacy-note{font-size:.76rem;color:#697386}.success-card{border:1px solid rgba(22,163,74,.20);background:rgba(240,253,244,.88);border-radius:14px;padding:.45rem .65rem}
.hr-hero{padding:.2rem 0 .35rem}.hr-badge{display:inline-block;padding:.28rem .7rem;border-radius:999px;background:rgba(255,255,255,.8);border:1px solid rgba(79,70,229,.16);font-weight:800;color:#4f46e5}.hr-card-title{font-weight:850}.hr-note{font-size:.86rem;color:#687083}.stButton>button[kind="secondary"]{font-weight:700!important}@media(max-width:900px){.block-container{padding-left:.55rem!important;padding-right:.55rem!important}h1{letter-spacing:-1.8px!important;font-size:2rem!important}h2{font-size:1.45rem!important}h3{font-size:1.15rem!important}.stButton>button,.stDownloadButton>button{min-height:44px!important}.stTextInput input,.stTextArea textarea{font-size:16px!important}.payment-cta{padding:.45rem}.privacy-note{font-size:.72rem}}
</style>
""", unsafe_allow_html=True)

DEFAULTS = {
    "document_ready": False, "interaction_id": None, "document_name": None,
    "document_type": None, "document_bytes": None, "summary": None,
    "messages": [], "active_view": "Summary", "extracted_information": None,
    "extracted_line_items": None, "analysis_result": None,
    "pricing_open": False, "home_choice": None, "document_tool_choice": "Summary", "hr_section": "Salary & HR Calculators",
    "resume_analysis": None, "resume_file_bytes": None, "resume_file_name": None, "resume_file_type": None,
    "resume_target_role": "", "improved_resume_bytes": None, "improved_resume_name": None,
    "resume_fix_paid_demo": False,
    "jd_text": "", "jd_analysis": None, "improved_jd_bytes": None, "improved_jd_name": None, "jd_fix_paid_demo": False,
    "offer_analysis": None, "offer_file_bytes": None, "offer_file_name": None, "offer_file_type": None,
    "improved_offer_bytes": None, "improved_offer_name": None, "offer_fix_paid_demo": False,
    "cover_letter": None, "cover_letter_bytes": None, "cover_letter_name": None, "cover_letter_paid_demo": False,
}
for key, value in DEFAULTS.items():
    if key not in st.session_state:
        st.session_state[key] = [] if isinstance(value, list) else value

@st.cache_resource
def get_client():
    try:
        api_key = st.secrets["GEMINI_API_KEY"]
    except Exception:
        st.error("GEMINI_API_KEY is missing from Streamlit Secrets.")
        st.stop()
    return genai.Client(api_key=api_key)


def auth_configured():
    try:
        return "auth" in st.secrets
    except Exception:
        return False

def logged_in():
    try:
        return bool(st.user.is_logged_in)
    except Exception:
        return False

def display_name():
    try:
        name = getattr(st.user, "name", None)
        email = getattr(st.user, "email", None)
        return str(name or (str(email).split("@")[0] if email else "User"))
    except Exception:
        return "User"

# PUBLIC-FIRST EXPERIENCE:
# Login is NOT required for the free Nexora workflow.
# Google authentication remains available for the paid checkout flow.
client = get_client()


def get_mime_type(file):
    if file.type:
        return file.type
    name = file.name.lower()
    if name.endswith(".pdf"): return "application/pdf"
    if name.endswith(".png"): return "image/png"
    if name.endswith(".jpg") or name.endswith(".jpeg"): return "image/jpeg"
    return "application/octet-stream"

def validate_file(file):
    mime = get_mime_type(file)
    if mime == "application/pdf" and file.size > MAX_PDF_SIZE:
        return False, "PDF files must be 50 MB or smaller."
    if mime.startswith("image/") and file.size > MAX_IMAGE_SIZE:
        return False, "Images must be 20 MB or smaller."
    if mime != "application/pdf" and not mime.startswith("image/"):
        return False, "Please upload a PDF, PNG, JPG or JPEG."
    return True, ""

def make_document_part(file):
    return {"type": "document" if get_mime_type(file) == "application/pdf" else "image", "data": base64.b64encode(file.getvalue()).decode("utf-8"), "mime_type": get_mime_type(file)}

def money(value):
    return f"₹{value:,.0f}"


def safe_float(value):
    try:
        return float(value)
    except Exception:
        return 0.0


def validate_ai_file(file):
    if file is None:
        return False, "Please upload a file."
    if file.size > MAX_AI_FILE_SIZE:
        return False, "Please keep the file below 20 MB for this first version."
    mime = file.type or ""
    if mime != "application/pdf" and not mime.startswith("image/"):
        return False, "Please upload a PDF, PNG, JPG or JPEG."
    return True, ""


def ai_text(prompt, file=None):
    """Run a short Gemini AI task, optionally with an uploaded PDF/image.

    The HR AI tools use the same Interactions API pattern as the main
    document workflow. This keeps PDF/image inputs inline and avoids the
    undefined ai_text() error that occurred in the Phase 2 build.
    """
    try:
        inputs = [{"type": "text", "text": prompt}]
        if file is not None:
            inputs.append(make_document_part(file))

        response = client.interactions.create(
            model=MODEL,
            input=inputs,
            store=False,
            generation_config={"thinking_level": "minimal"},
        )
        result = getattr(response, "output_text", "") or ""
        if not result:
            st.error("Nexora did not receive a usable AI response.")
            return ""
        return result
    except Exception as error:
        text = str(error)
        st.error("Nexora could not complete the AI analysis.")
        if "429" in text:
            st.warning("The Gemini request limit was reached. Please wait a moment and try again.")
        elif "500" in text or "503" in text:
            st.warning("Gemini is temporarily busy. Please try again.")
        elif "400" in text:
            st.warning("The uploaded file or AI request could not be processed. Please try a different PDF/image.")
        else:
            st.warning("Please try the analysis again.")
        logger.exception("Nexora AI request failed: %s", text)
        return ""


def _set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_border(cell, color="D9E2F3", size="6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def _set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def _set_doc_margins(section):
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)


def _add_field(paragraph, instruction):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = instruction
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def _style_document(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(42, 52, 72)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in [
        ("Title", 24, "173B6C"),
        ("Heading 1", 13.5, "173B6C"),
        ("Heading 2", 11.5, "2563A6"),
    ]:
        stl = styles[name]
        stl.font.name = "Aptos Display" if name == "Title" else "Aptos"
        stl.font.size = Pt(size)
        stl.font.bold = True
        stl.font.color.rgb = RGBColor.from_string(color)
        stl.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        stl.paragraph_format.space_after = Pt(5)
    for sec in doc.sections:
        _set_doc_margins(sec)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = "CONFIDENTIAL  •  For authorized HR / business use only  •  Page "
        _add_field(footer, "PAGE")
        for r in footer.runs:
            r.font.name = "Aptos"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(105, 116, 135)


def _add_header(doc, title, subtitle):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(4.9)
    table.columns[1].width = Inches(1.9)
    left, right = table.rows[0].cells
    _set_cell_shading(left, "F4F8FD")
    _set_cell_shading(right, "EAF2FF")
    _set_cell_border(left, "D8E5F5", "5")
    _set_cell_border(right, "D8E5F5", "5")
    left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("[COMPANY NAME]")
    r.bold = True; r.font.name = "Aptos Display"; r.font.size = Pt(16); r.font.color.rgb = RGBColor(23,59,108)
    p2 = left.add_paragraph("Human Resources / People Operations")
    p2.paragraph_format.space_after = Pt(0)
    for r in p2.runs:
        r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(91,105,125)
    p3 = right.paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p3.add_run("HR DOCUMENT")
    r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor(37,99,235)
    p4 = right.add_paragraph(subtitle)
    p4.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p4.paragraph_format.space_after = Pt(0)
    for r in p4.runs:
        r.font.size = Pt(8); r.font.color.rgb = RGBColor(91,105,125)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    title_p = doc.add_paragraph(style="Title")
    title_p.add_run(title)
    title_p.paragraph_format.space_after = Pt(2)
    accent = doc.add_paragraph()
    accent.paragraph_format.space_after = Pt(8)
    run = accent.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    run.font.size = Pt(5); run.font.color.rgb = RGBColor(37,99,235)


def _add_metadata_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    for i, (label, value) in enumerate(rows):
        table.columns[0].width = Inches(1.55)
        table.columns[1].width = Inches(5.25)
        c1, c2 = table.rows[i].cells
        _set_cell_shading(c1, "F7F9FC")
        _set_cell_border(c1); _set_cell_border(c2)
        p = c1.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label.upper()); r.bold=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor(80,96,120)
        p = c2.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value); r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(42,52,72)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_section(doc, heading, lines):
    doc.add_heading(heading, level=1)
    for line in lines:
        p = doc.add_paragraph()
        if line.startswith("• "):
            p.style = doc.styles["Normal"]
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.15)
        p.add_run(line)


def _add_signature(doc, signatory="[AUTHORIZED SIGNATORY]"):
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(3.35); table.columns[1].width = Inches(3.35)
    for cell, heading in zip(table.rows[0].cells, ["For [COMPANY NAME]", "Employee / Recipient"]):
        _set_cell_border(cell, "FFFFFF", "0")
        p = cell.paragraphs[0]; p.add_run("\n\n____________________________").bold=True
        p.add_run(f"\n{heading}\n[NAME / SIGNATURE]\n[DATE]")
        for r in p.runs: r.font.size=Pt(9); r.font.color.rgb=RGBColor(70,82,103)


def docx_bytes(title, subtitle, sections, metadata=None, form=False):
    doc = Document()
    _style_document(doc)
    _add_header(doc, title, subtitle)
    if metadata:
        _add_metadata_table(doc, metadata)
    for heading, body in sections:
        _add_section(doc, heading, body)
    if form:
        table = doc.add_table(rows=1, cols=3)
        table.autofit = False
        headers = ["Area", "Response / Rating", "Comments"]
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            _set_cell_shading(cell, "173B6C"); _set_cell_border(cell, "173B6C")
            p = cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(h); r.bold=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(255,255,255)
        _set_repeat_table_header(table.rows[0])
        for area in ["Overall experience", "Manager support", "Work environment", "Growth / career", "Compensation", "Reason for leaving", "Would recommend employer"]:
            cells=table.add_row().cells
            for c in cells: _set_cell_border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[0].text=area; cells[1].text="[ ] 1  [ ] 2  [ ] 3  [ ] 4  [ ] 5"; cells[2].text="[COMMENTS]"
            for c in cells:
                for r in c.paragraphs[0].runs: r.font.size=Pt(8.5)
        doc.add_paragraph("Additional comments: ________________________________________________________________")
        doc.add_paragraph("________________________________________________________________________________")
    else:
        _add_signature(doc)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _resume_clean_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


def _resume_json_from_ai(text):
    """Extract JSON from the model while tolerating a fenced JSON response."""
    cleaned = _resume_clean_text(text)
    try:
        return json.loads(cleaned)
    except Exception:
        match = re.search(r"\{.*\}", cleaned, flags=re.S)
        if match:
            try:
                return json.loads(match.group(0))
            except Exception:
                return None
    return None


def _docx_add_resume_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(str(text).upper())
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(37, 99, 235)
    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(4)
    rr = line.add_run("────────────────────────────────────────────────────────")
    rr.font.size = Pt(4.5)
    rr.font.color.rgb = RGBColor(191, 219, 254)


def _docx_add_resume_bullets(doc, values):
    if not isinstance(values, list):
        return
    for value in values:
        if not str(value).strip():
            continue
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        r = p.add_run(str(value).strip())
        r.font.name = "Aptos"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(45, 55, 72)


def _docx_add_resume_entries(doc, entries):
    if not isinstance(entries, list):
        return
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        title = str(entry.get("title", "")).strip()
        organization = str(entry.get("organization", "")).strip()
        dates = str(entry.get("dates", "")).strip()
        location = str(entry.get("location", "")).strip()
        if title:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(title)
            r.bold = True; r.font.name = "Aptos Display"; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(37, 45, 63)
            meta = " • ".join([x for x in [organization, location, dates] if x])
            if meta:
                r2 = p.add_run("  " + meta)
                r2.italic = True; r2.font.name = "Aptos"; r2.font.size = Pt(8.8); r2.font.color.rgb = RGBColor(91, 105, 125)
        _docx_add_resume_bullets(doc, entry.get("bullets", []))


def professional_resume_docx(data):
    """Build a clean ATS-friendly resume without Nexora branding in the file."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    _style_document(doc)

    name = str(data.get("name", "PROFESSIONAL RESUME")).strip() or "PROFESSIONAL RESUME"
    role = str(data.get("target_title", "")).strip()
    contact = [str(data.get(k, "")).strip() for k in ["phone", "email", "location", "linkedin"]]
    contact = [x for x in contact if x]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(name)
    r.bold = True; r.font.name = "Aptos Display"; r.font.size = Pt(21); r.font.color.rgb = RGBColor(23, 59, 108)
    if role:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
        r = p.add_run(role); r.bold=True; r.font.name="Aptos"; r.font.size=Pt(11); r.font.color.rgb=RGBColor(37,99,235)
    if contact:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
        r = p.add_run("  |  ".join(contact)); r.font.name="Aptos"; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(91,105,125)

    summary = str(data.get("summary", "")).strip()
    if summary:
        _docx_add_resume_heading(doc, "Professional Summary")
        p = doc.add_paragraph(summary); p.paragraph_format.space_after = Pt(3)
        for r in p.runs: r.font.name="Aptos"; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(45,55,72)

    skills = data.get("skills", [])
    if isinstance(skills, list) and skills:
        _docx_add_resume_heading(doc, "Core Skills")
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        r = p.add_run("  •  ".join([str(x).strip() for x in skills if str(x).strip()]))
        r.font.name="Aptos"; r.font.size=Pt(9.2); r.font.color.rgb=RGBColor(45,55,72)

    if data.get("experience"):
        _docx_add_resume_heading(doc, "Professional Experience")
        _docx_add_resume_entries(doc, data.get("experience"))
    if data.get("projects"):
        _docx_add_resume_heading(doc, "Selected Projects")
        _docx_add_resume_entries(doc, data.get("projects"))
    if data.get("education"):
        _docx_add_resume_heading(doc, "Education")
        _docx_add_resume_entries(doc, data.get("education"))
    if data.get("certifications"):
        _docx_add_resume_heading(doc, "Certifications")
        _docx_add_resume_bullets(doc, data.get("certifications"))
    if data.get("achievements"):
        _docx_add_resume_heading(doc, "Achievements")
        _docx_add_resume_bullets(doc, data.get("achievements"))
    if data.get("additional"):
        _docx_add_resume_heading(doc, "Additional Information")
        _docx_add_resume_bullets(doc, data.get("additional"))

    # Neutral resume footer; deliberately no Nexora promotional text.
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Professional Resume  •  Candidate-provided information only")
    fr.font.name="Aptos"; fr.font.size=Pt(7.5); fr.font.color.rgb=RGBColor(120,130,145)

    output = io.BytesIO(); doc.save(output); return output.getvalue()


def improve_resume(resume_file, target_role, analysis, resume_file_type="application/pdf", resume_file_name="resume.pdf"):
    prompt = f"""
You are Nexora's senior executive resume writer and ATS optimization specialist.
Create a highly professional, ATS-friendly resume from the uploaded candidate resume.
Target role: {target_role or 'Not specified'}

Existing resume analysis:
{analysis or 'Not available'}

Your task is to FIX the weaknesses identified in the analysis while preserving factual accuracy.
You may rewrite weak wording, improve structure, remove repetition, strengthen action language,
make achievements clearer, improve keyword alignment, and create a concise professional summary.
You MUST NOT invent employers, job titles, dates, degrees, certifications, technologies, metrics,
responsibilities, awards or achievements. Never create a number that was not present in the source.
If a useful field is absent, leave it empty rather than guessing.
Keep the candidate's actual career history intact.

Return ONLY valid JSON in this exact shape:
{{
  "name":"",
  "target_title":"",
  "phone":"",
  "email":"",
  "location":"",
  "linkedin":"",
  "summary":"",
  "skills":[""],
  "experience":[{{"title":"","organization":"","dates":"","location":"","bullets":[""]}}],
  "projects":[{{"title":"","organization":"","dates":"","location":"","bullets":[""]}}],
  "education":[{{"title":"","organization":"","dates":"","location":"","bullets":[""]}}],
  "certifications":[""],
  "achievements":[""],
  "additional":[""]
}}

Use only information found in the uploaded resume. Make the finished resume concise, modern,
professional and ATS-readable. Rewrite bullet points into strong action-oriented language when
supported by the source. Remove decorative symbols, tables, columns, photos and unnecessary
personal details. Do not include a cover letter. Do not mention Nexora in the resume content.
"""
    # Rebuild a lightweight upload object from session-persisted bytes so the
    # improvement step still works after Streamlit reruns.
    if isinstance(resume_file, (bytes, bytearray)):
        upload = io.BytesIO(bytes(resume_file))
        upload.name = resume_file_name or "resume.pdf"
        upload.type = resume_file_type or "application/pdf"
        resume_file = upload
    raw = ai_text(prompt, resume_file)
    if not raw:
        return None, None
    data = _resume_json_from_ai(raw)
    if not isinstance(data, dict):
        st.error("Nexora received an unexpected resume format. Please try again.")
        return None, None
    try:
        output = professional_resume_docx(data)
    except Exception as error:
        st.error("The improved resume could not be formatted.")
        logger.exception("Nexora operation failed: %s", error)
        return None, None
    safe_name = re.sub(r"[^A-Za-z0-9]+", "_", str(data.get("name") or "professional_resume")).strip("_") or "professional_resume"
    return output, f"{safe_name}_professional_resume.docx"



def professional_simple_docx(title, subtitle, sections, filename_base="professional_document"):
    """Create a polished, neutral Word document for paid-demo career outputs."""
    doc = Document()
    _style_document(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    _add_header(doc, title, subtitle)
    for heading, body in sections:
        _add_section(doc, heading, body)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Professional document  •  Prepared from user-provided information")
    fr.font.name = "Aptos"; fr.font.size = Pt(7.5); fr.font.color.rgb = RGBColor(120,130,145)
    out = io.BytesIO(); doc.save(out)
    safe = re.sub(r"[^A-Za-z0-9]+", "_", filename_base).strip("_") or "professional_document"
    return out.getvalue(), f"{safe}.docx"


def improve_job_description(jd, analysis, target_role=""):
    prompt = f"""
You are a senior HR business partner and recruitment copywriter.
Rewrite the supplied job description into a highly professional, clear, inclusive and ATS-friendly job description.
Target role: {target_role or 'Use the role stated in the source'}
Existing analysis:
{analysis or 'Not available'}

Rules:
- Preserve every factual requirement from the source where possible.
- Do not invent salary, benefits, location, company facts, reporting lines, qualifications or experience requirements.
- Remove repetition, vague filler, discriminatory wording and unnecessary jargon.
- Make responsibilities action-oriented and scannable.
- Separate must-have and nice-to-have requirements.
- Add a concise role overview and success profile only when supported by the source; otherwise keep them neutral.
- Do not mention Nexora.

Return ONLY valid JSON:
{{
  "title":"",
  "overview":"",
  "responsibilities":[""],
  "must_have":[""],
  "nice_to_have":[""],
  "skills":[""],
  "education":[""],
  "experience":[""],
  "benefits":[""],
  "application_notes":[""]
}}

SOURCE JOB DESCRIPTION:
---
{jd}
---
"""
    raw = ai_text(prompt)
    data = _resume_json_from_ai(raw) if raw else None
    if not isinstance(data, dict):
        st.error("Nexora could not create the improved job description. Please try again.")
        return None, None
    sections=[]
    for key, heading in [("overview","Role Overview"),("responsibilities","Key Responsibilities"),("must_have","Must-Have Requirements"),("nice_to_have","Nice-to-Have Requirements"),("skills","Skills & Competencies"),("education","Education"),("experience","Experience"),("benefits","Benefits"),("application_notes","Application Notes")]:
        val=data.get(key)
        if isinstance(val,list): body=[str(x) for x in val if str(x).strip()]
        else: body=[str(val).strip()] if str(val).strip() else []
        if body: sections.append((heading, body))
    return professional_simple_docx(data.get("title") or target_role or "Professional Job Description", "Refined recruitment-ready version", sections, "professional_job_description")


def generate_cover_letter(candidate_resume, jd, target_role=""):
    prompt=f"""
You are an expert executive career writer.
Create a highly professional, tailored cover letter using ONLY the candidate information in the resume and the requirements in the job description.
Target role: {target_role or 'Use the role from the job description'}

Rules:
- Never invent employers, achievements, metrics, qualifications, dates or experience.
- Do not claim skills that are not supported by the resume.
- Match the candidate's real experience to the job requirements where evidence exists.
- Keep it concise, confident and natural, approximately 350-500 words.
- Use placeholders such as [Hiring Manager Name] only where information is genuinely unavailable.
- Do not mention Nexora or AI.
- Return ONLY valid JSON: {{"candidate_name":"","subject":"","salutation":"","body_paragraphs":[""],"closing":""}}

CANDIDATE RESUME:
---
{candidate_resume}
---
JOB DESCRIPTION:
---
{jd}
---
"""
    raw=ai_text(prompt)
    data=_resume_json_from_ai(raw) if raw else None
    if not isinstance(data,dict):
        st.error("Nexora could not create the cover letter. Please try again.")
        return None,None
    doc=Document(); _style_document(doc)
    section=doc.sections[0]; section.top_margin=Inches(.7); section.bottom_margin=Inches(.7); section.left_margin=Inches(.8); section.right_margin=Inches(.8)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("COVER LETTER"); r.bold=True; r.font.name="Aptos Display"; r.font.size=Pt(19); r.font.color.rgb=RGBColor(23,59,108)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(14); r=p.add_run(str(data.get("subject","")).strip()); r.bold=True; r.font.name="Aptos"; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor(37,99,235)
    p=doc.add_paragraph(str(data.get("salutation") or "Dear Hiring Manager,")); p.paragraph_format.space_after=Pt(9)
    for para in data.get("body_paragraphs",[]):
        if str(para).strip():
            p=doc.add_paragraph(str(para).strip()); p.paragraph_format.space_after=Pt(9)
    p=doc.add_paragraph(str(data.get("closing") or "Sincerely,")); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(22)
    candidate_name = str(data.get("candidate_name") or "Candidate Name").strip()
    p.add_run(candidate_name).bold=True
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; fr=fp.add_run("Professional cover letter  •  Prepared from user-provided information"); fr.font.name="Aptos"; fr.font.size=Pt(7.5); fr.font.color.rgb=RGBColor(120,130,145)
    out=io.BytesIO(); doc.save(out); return out.getvalue(), "tailored_cover_letter.docx"


def create_offer_review_pack(offer_file, analysis):
    prompt=f"""
You are a senior HR document specialist.
Based on the uploaded offer letter and the analysis below, create a professional candidate decision-support pack.
Do not provide legal advice and do not invent facts.
Analysis:
{analysis or 'Not available'}
Return ONLY valid JSON:
{{"title":"Offer Letter Review Pack","executive_summary":"","key_terms":[""],"attention_items":[""],"questions_for_hr":[""],"before_accepting":[""]}}
"""
    raw=ai_text(prompt, offer_file); data=_resume_json_from_ai(raw) if raw else None
    if not isinstance(data,dict):
        st.error("Nexora could not create the offer review pack. Please try again.")
        return None,None
    sections=[]
    for key,heading in [("executive_summary","Executive Summary"),("key_terms","Key Terms"),("attention_items","Items Requiring Attention"),("questions_for_hr","Questions for HR"),("before_accepting","Before Accepting")]:
        val=data.get(key); body=val if isinstance(val,list) else [str(val)]
        body=[str(x) for x in body if str(x).strip()]
        if body: sections.append((heading,body))
    return professional_simple_docx("Offer Letter Review Pack","Candidate decision-support document",sections,"offer_letter_review_pack")

def reset_workspace():
    for key, value in DEFAULTS.items():
        st.session_state[key] = [] if isinstance(value, list) else value

def analyze_document(file):
    prompt = """
You are Nexora, a professional AI document assistant.
Read and understand the uploaded document carefully.
Return:
## Executive Summary
Give a concise explanation of what the document is about.
## Key Information
List the most important facts including names, dates, amounts, organizations, reference numbers, addresses, important terms and other critical information.
## Key Takeaways
Give the most useful points a person should know.
## Risks and Concerns
Identify missing information, potential risks, inconsistencies, unusual clauses and items requiring verification. If none are obvious, say so.
## Suggested Questions
Give 5 useful questions the user could ask about this document.
Rules: Do not invent information. Use only information present in the document. Preserve dates and numbers accurately. If unclear, say so. Keep the result professional and easy to scan.
"""
    return client.interactions.create(model=MODEL, input=[{"type":"text","text":prompt}, make_document_part(file)], store=True, generation_config={"thinking_level":"minimal"})

def process_document(file):
    try:
        with st.spinner("Nexora is understanding your document..."):
            interaction = analyze_document(file)
        interaction_id = getattr(interaction, "id", None)
        output = getattr(interaction, "output_text", "")
        if not interaction_id or not output:
            st.error("Nexora did not receive a valid Gemini response.")
            return False
        st.session_state.interaction_id = interaction_id
        st.session_state.summary = output
        st.session_state.document_name = file.name
        st.session_state.document_type = get_mime_type(file)
        st.session_state.document_bytes = file.getvalue()
        st.session_state.document_ready = True
        st.session_state.messages = []
        st.session_state.extracted_information = None
        st.session_state.extracted_line_items = None
        st.session_state.analysis_result = None
        st.session_state.active_view = "Summary"
        return True
    except Exception as error:
        text = str(error)
        st.error("Nexora could not analyze the document.")
        if "429" in text: st.warning("The Gemini request limit was reached. Please wait and try again.")
        elif "500" in text or "503" in text: st.warning("Gemini is temporarily busy. Please try again.")
        else: st.warning("Please try the Analyze Document button again.")
        logger.exception("Nexora operation failed: %s", text)
        return False

def ask_document(question, placeholder):
    try:
        stream = client.interactions.create(model=MODEL, previous_interaction_id=st.session_state.interaction_id, input=question, store=True, stream=True, generation_config={"thinking_level":"minimal"})
        answer, latest_id = "", None
        for event in stream:
            event_type = getattr(event, "event_type", None)
            if event_type == "step.delta":
                delta = getattr(event, "delta", None)
                if delta and getattr(delta, "type", None) == "text":
                    text = getattr(delta, "text", "")
                    if text:
                        answer += text
                        placeholder.markdown(answer)
            elif event_type == "interaction.complete":
                latest_id = getattr(event, "id", None)
        if latest_id: st.session_state.interaction_id = latest_id
        return answer
    except Exception as error:
        st.error("Nexora could not answer the question.")
        logger.exception("Nexora operation failed: %s", error)
        return ""

def extract_data():
    prompt = """
Extract structured information from the document. Return ONLY valid JSON:
{"document_information":[{"field":"","value":""}],"line_items":[{"description":"","quantity":"","unit_price":"","amount":""}]}
Rules: extract only information present; never invent; use empty strings when unavailable; preserve dates and numbers; valid JSON only.
"""
    try:
        with st.spinner("Extracting structured information..."):
            response = client.interactions.create(model=MODEL, previous_interaction_id=st.session_state.interaction_id, input=prompt, store=True, generation_config={"thinking_level":"minimal"})
        st.session_state.interaction_id = response.id
        text = getattr(response, "output_text", "")
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if not match:
            st.error("Nexora could not extract structured data.")
            return
        data = json.loads(match.group())
        st.session_state.extracted_information = data.get("document_information", [])
        st.session_state.extracted_line_items = data.get("line_items", [])
        st.session_state.active_view = "Extract Data"
        st.rerun()
    except Exception as error:
        st.error("Data extraction failed.")
        logger.exception("Nexora operation failed: %s", error)

def run_deep_analysis():
    prompt = """
Perform a detailed analysis of this document. Focus on important risks, missing information, dates, amounts, unusual clauses, inconsistencies, information requiring human attention and practical next steps. Do not invent anything. Clearly distinguish facts from observations. Use clear headings and bullet points.
"""
    try:
        with st.spinner("Nexora is performing deeper analysis..."):
            response = client.interactions.create(model=MODEL, previous_interaction_id=st.session_state.interaction_id, input=prompt, store=True, generation_config={"thinking_level":"minimal"})
        st.session_state.interaction_id = response.id
        st.session_state.analysis_result = getattr(response, "output_text", "")
        st.session_state.active_view = "Analyze"
        st.rerun()
    except Exception as error:
        st.error("Deep analysis failed.")
        logger.exception("Nexora operation failed: %s", error)

def create_excel():
    info_df = pd.DataFrame(st.session_state.extracted_information or [])
    items_df = pd.DataFrame(st.session_state.extracted_line_items or [])
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        if not info_df.empty: info_df.to_excel(writer, sheet_name="Document Information", index=False)
        if not items_df.empty: items_df.to_excel(writer, sheet_name="Line Items", index=False)
    return output.getvalue()


# ============================================================
# NEXORA PRODUCT NAVIGATION
# ============================================================
if "main_section" not in st.session_state:
    st.session_state.main_section = "Home"

logo, home_nav, pricing_nav, account = st.columns([2.7, 1.0, 1.0, 1.35], gap="small", vertical_alignment="center")

with logo:
    st.markdown("# ✦ NEXORA")
    st.caption("AI WORK • DOCUMENTS • CAREER")

with home_nav:
    if st.button("⌂ Home", use_container_width=True, key="nav_home"):
        st.session_state.main_section = "Home"
        st.session_state.pricing_open = False
        st.rerun()

with pricing_nav:
    if st.button("Pricing", use_container_width=True, key="nav_pricing"):
        st.session_state.pricing_open = not st.session_state.pricing_open
        st.rerun()

with account:
    if logged_in():
        a, b = st.columns([2.4, 1], gap="small")
        with a:
            st.caption(f"● {display_name()}")
        with b:
            if st.button("↪", help="Log out", key="nav_logout"):
                st.logout()
    else:
        st.caption("Free access")

st.divider()

# ============================================================
# PRICING / PAYMENT LOGIN GATE
# ============================================================
if st.session_state.pricing_open:
    st.markdown("# ✦ Nexora Plans")
    st.caption("Explore Nexora freely. Login is requested only when you choose a paid plan.")

    free_col, pro_col = st.columns(2, gap="small", vertical_alignment="top")

    with free_col:
        with st.container(border=True):
            st.markdown("## 🆓 Free")
            st.markdown("### Explore Nexora")
            st.caption("No account required for the current free workflows.")
            for item in [
                "AI document summary",
                "Document Q&A",
                "Structured data extraction",
                "HR salary calculators",
                "Resume / JD / offer analysis",
                "Starter HR templates",
            ]:
                st.write(f"✓ {item}")
            st.button("Current plan", disabled=True, use_container_width=True, key="free_plan")

    with pro_col:
        with st.container(border=True):
            st.markdown("## ✦ Nexora Pro")
            st.markdown("### Paid plan")
            st.caption("Premium usage and future account-based features.")
            for item in [
                "Higher usage limits",
                "Premium document workflows",
                "Detailed HR & career reports",
                "Account-linked purchase",
                "Saved-document and calculation history",
            ]:
                st.write(f"✓ {item}")

            if not logged_in():
                st.info("Sign in is required only when you start checkout.")
                if st.button(
                    "G  Login & Continue to Checkout",
                    type="primary",
                    use_container_width=True,
                    key="pricing_login",
                ):
                    if auth_configured():
                        st.login()
                    else:
                        st.warning("Google login is not configured yet.")
            else:
                st.success(f"Signed in as {display_name()}.")
                st.button(
                    "Continue to Payment",
                    type="primary",
                    use_container_width=True,
                    key="continue_payment",
                )
                st.caption(
                    "Google Login is ready. The payment gateway can be connected "
                    "to this button in the next phase."
                )

    st.divider()

# ============================================================
# HR & CAREER HUB
# ============================================================
def render_hr_hub():
    st.caption("✦ NEXORA HR & CAREER • SALARY • CAREER • TOOLS")
    st.markdown("# HR work, career decisions. ✦ **Made simpler.**")
    st.write("Salary calculators, AI career tools and practical HR templates — built into Nexora.")

    if "hr_section" not in st.session_state:
        st.session_state.hr_section = "Salary & HR Calculators"
    st.markdown("### Choose a Career Workspace")
    hr_a, hr_b, hr_c = st.columns(3, gap="small")
    hr_options = [
        (hr_a, "💰  Salary & HR Calculators", "Salary, increment, gratuity and notice-period tools.", "Salary & HR Calculators"),
        (hr_b, "🤖  AI Career Tools", "Resume, JD, offer-letter and cover-letter tools.", "AI Career Tools"),
        (hr_c, "📄  HR Templates", "Professional editable workplace documents.", "HR Templates"),
    ]
    for col, label, desc, value in hr_options:
        with col:
            if st.button(label, use_container_width=True, key=f"hr_tab_{value}"):
                st.session_state.hr_section = value
                st.rerun()
            st.caption(desc)
    section = st.session_state.hr_section

    if section == "Salary & HR Calculators":
        st.markdown("## 💰 Salary & HR Calculators")
        st.caption("Free calculators designed for quick, practical answers.")
        calculator = st.selectbox("Choose a calculator", ["CTC → Take-Home Salary","Salary Increment Calculator","Gratuity Estimator","Notice Period Salary Calculator"])
        if calculator == "CTC → Take-Home Salary":
            left,right=st.columns([1.1,1],gap="small",vertical_alignment="top")
            with left:
                with st.container(border=True):
                    st.subheader("🧾 Enter your salary structure")
                    annual_ctc=st.number_input("Annual CTC (₹)",min_value=0.0,value=600000.0,step=10000.0)
                    basic_pct=st.number_input("Basic salary as % of CTC",min_value=0.0,max_value=100.0,value=40.0,step=1.0)
                    variable_pct=st.number_input("Variable / bonus as % of CTC",min_value=0.0,max_value=100.0,value=0.0,step=1.0)
                    employee_pf_rate=st.number_input("Employee PF rate on Basic (%)",min_value=0.0,max_value=20.0,value=12.0,step=0.5)
                    monthly_pt=st.number_input("Professional tax / other fixed monthly deduction (₹)",min_value=0.0,value=200.0,step=50.0)
                    other_monthly=st.number_input("Other monthly deductions (₹)",min_value=0.0,value=0.0,step=100.0)
                    pf_cap=st.number_input("PF wage base cap (₹/month, 0 = no cap)",min_value=0.0,value=15000.0,step=1000.0)
            basic_annual=annual_ctc*basic_pct/100; variable_annual=annual_ctc*variable_pct/100; fixed_annual=max(annual_ctc-variable_annual,0); basic_monthly=basic_annual/12; pf_base=min(basic_monthly,pf_cap) if pf_cap>0 else basic_monthly; employee_pf=pf_base*employee_pf_rate/100; monthly_gross=fixed_annual/12; take_home=max(monthly_gross-employee_pf-monthly_pt-other_monthly,0)
            with right:
                with st.container(border=True):
                    st.subheader("📊 Estimated result"); m1,m2=st.columns(2)
                    with m1: st.metric("Monthly gross",money(monthly_gross)); st.metric("Employee PF",money(employee_pf))
                    with m2: st.metric("Estimated take-home",money(take_home)); st.metric("Annual fixed pay",money(fixed_annual))
                    st.divider(); st.write(f"**Basic:** {money(basic_monthly)}/month"); st.write(f"**Variable / bonus:** {money(variable_annual)}/year"); st.write(f"**Other fixed deductions:** {money(monthly_pt+other_monthly)}/month")
                    st.info("This is an estimate, not a payroll statement. Actual take-home can differ based on tax regime, employer policy, PF treatment, insurance, professional tax and other deductions.")
        elif calculator == "Salary Increment Calculator":
            c1,c2=st.columns(2,gap="small")
            with c1:
                with st.container(border=True):
                    st.subheader("📈 Current salary"); current=st.number_input("Current annual CTC (₹)",min_value=0.0,value=600000.0,step=10000.0); increment=st.number_input("Increment (%)",min_value=-100.0,max_value=500.0,value=10.0,step=0.5); current_monthly=current/12
            with c2:
                with st.container(border=True):
                    new_ctc=current*(1+increment/100); increase=new_ctc-current; st.subheader("🚀 New salary"); st.metric("New annual CTC",money(new_ctc)); st.metric("Annual increase",money(increase)); st.metric("New monthly CTC",money(new_ctc/12)); st.caption(f"Current monthly CTC: {money(current_monthly)}")
            rows=[]
            for pct in [5,8,10,12,15,20,25,30]:
                nv=current*(1+pct/100); rows.append({"Increment":f"{pct}%","New CTC":money(nv),"Monthly CTC":money(nv/12),"Annual increase":money(nv-current)})
            st.dataframe(pd.DataFrame(rows),use_container_width=True,hide_index=True)
        elif calculator == "Gratuity Estimator":
            c1,c2=st.columns([1.05,1],gap="small")
            with c1:
                with st.container(border=True):
                    st.subheader("🏆 Gratuity inputs"); last_wage=st.number_input("Last drawn monthly wages used for calculation (₹)",min_value=0.0,value=30000.0,step=1000.0); years=st.number_input("Completed years of service",min_value=0,max_value=60,value=5,step=1); extra_months=st.number_input("Extra months",min_value=0,max_value=11,value=0,step=1); apply_rounding=st.checkbox("Treat 6+ extra months as one additional year",value=True); cap=st.number_input("Optional gratuity cap (₹, 0 = no cap)",min_value=0.0,value=0.0,step=10000.0)
            effective_years=years+(1 if apply_rounding and extra_months>=6 else 0); raw_gratuity=last_wage*15/26*effective_years; final_gratuity=min(raw_gratuity,cap) if cap>0 else raw_gratuity
            with c2:
                with st.container(border=True): st.subheader("💡 Estimated gratuity"); st.metric("Estimated amount",money(final_gratuity)); st.write(f"**Formula:** {money(last_wage)} × 15 ÷ 26 × {effective_years} year(s)"); st.info("This is an estimate. Eligibility and calculation can depend on the applicable law, employee category and service conditions. Verify the applicable rule before relying on the result.")
        else:
            c1,c2=st.columns(2,gap="small")
            with c1:
                with st.container(border=True): st.subheader("📅 Notice period"); monthly_salary=st.number_input("Monthly salary (₹)",min_value=0.0,value=50000.0,step=1000.0); notice_days=st.number_input("Notice period (days)",min_value=0,max_value=365,value=30,step=1); working_days=st.number_input("Working days used for daily-rate estimate",min_value=1,max_value=31,value=30,step=1)
            with c2:
                with st.container(border=True): daily_rate=monthly_salary/working_days if working_days else 0; estimated=daily_rate*notice_days; st.subheader("💼 Estimate"); st.metric("Daily salary basis",money(daily_rate)); st.metric("Notice-period salary",money(estimated)); st.info("Notice pay depends on the employment contract, applicable law and company policy. This is a simple estimate.")

    elif section == "AI Career Tools":
        st.markdown("## 🤖 AI Career Tools")
        st.caption("Free analysis gives the user value first. Paid-demo actions turn the identified work into a finished professional document.")
        st.caption("🔒 Career files may contain personal or employment information. Upload only documents you are authorized to process.")
        tool=st.selectbox("Choose an AI career tool",["Resume Analyzer","Job Description Analyzer","Offer Letter Analyzer","Cover Letter Generator"])

        if tool == "Resume Analyzer":
            st.subheader("📄 Resume Analyzer")
            resume=st.file_uploader("Upload resume",type=["pdf","png","jpg","jpeg"],key="resume_upload")
            target_role=st.text_input("Target role (optional)",placeholder="e.g. HR Manager, Talent Acquisition Specialist",key="resume_target_role_input")
            if resume:
                valid,msg=validate_ai_file(resume)
                if not valid: st.error(msg)
                elif st.button("✦ Analyze Resume",type="primary",use_container_width=True,key="analyze_resume_btn"):
                    prompt=f"""
You are Nexora HR, an expert Indian recruitment and career assistant.
Analyze the uploaded resume carefully. Target role: {target_role or 'Not specified'}.
Return a practical report with these sections:
1. Overall Resume Score / 100
2. Executive Assessment
3. ATS Readiness
4. Strongest Skills and Evidence
5. Missing or Weak Areas
6. Experience & Achievement Quality
7. Keywords to Add
8. Formatting / Clarity Issues
9. Role Match Assessment
10. Top 10 Changes to Improve Interview Chances
11. Improved Professional Summary
12. 5 Suggested Interview Questions
Do not invent experience, qualifications or achievements. If information is missing, say so.
"""
                    with st.spinner("Nexora is analyzing the resume..."): result=ai_text(prompt,resume)
                    if result:
                        st.session_state.resume_analysis=result; st.session_state.resume_file_bytes=resume.getvalue(); st.session_state.resume_file_name=resume.name; st.session_state.resume_file_type=get_mime_type(resume); st.session_state.resume_target_role=target_role; st.session_state.improved_resume_bytes=None; st.session_state.improved_resume_name=None; st.session_state.resume_fix_paid_demo=False
            if st.session_state.resume_analysis:
                # HIGH-VISIBILITY PAID ACTION: placed before the long report so the user sees the next step immediately.
                with st.container(border=True):
                    st.markdown("### 🚀 Fix my resume — one-time ₹99")
                    pcol,scol=st.columns([1.0,2.0],vertical_alignment="center")
                    with pcol:
                        if st.button("💳 Pay ₹99 (Temporary)",type="primary",use_container_width=True,key="resume_pay_demo_top"):
                            with st.spinner("Demo payment confirmed. Nexora is rebuilding your resume..."):
                                b,n=improve_resume(st.session_state.resume_file_bytes,st.session_state.resume_target_role,st.session_state.resume_analysis,st.session_state.resume_file_type,st.session_state.resume_file_name)
                            if b: st.session_state.improved_resume_bytes=b; st.session_state.improved_resume_name=n; st.session_state.resume_fix_paid_demo=True
                    with scol:
                        st.write("AI fixes the weaknesses identified in your analysis, improves ATS readability and creates a professional editable resume.")
                        st.caption("Temporary demo checkout • No real payment is processed")
                st.divider(); st.subheader("📊 Resume Assessment"); st.markdown(st.session_state.resume_analysis)
                if st.session_state.improved_resume_bytes:
                    st.success("Demo payment successful — your improved professional resume is ready.")
                    d1,d2=st.columns([1,1],gap="small")
                    with d1:
                        st.download_button("⬇️ Download Professional Resume (Word)",data=st.session_state.improved_resume_bytes,file_name=st.session_state.improved_resume_name or "professional_resume.docx",mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True,key="download_improved_resume")
                    with d2:
                        st.caption("No Nexora promotional branding is inserted into the downloaded resume.")

        elif tool == "Job Description Analyzer":
            st.subheader("🎯 Job Description Analyzer")
            st.write("Analyze a JD for free, then optionally purchase a professionally rewritten recruitment-ready version.")
            jd=st.text_area("Paste the job description",height=300,placeholder="Paste the complete job description here...",key="jd_text_area")
            resume_context=st.text_area("Optional: paste your resume/profile summary",height=130,placeholder="Useful for candidate-to-JD match analysis...",key="jd_resume_context")
            if st.button("✦ Analyze Job Description",type="primary",use_container_width=True,key="analyze_jd_btn"):
                if len(jd.strip())<40: st.warning("Please paste a fuller job description so the analysis is useful.")
                else:
                    prompt=f"""
You are Nexora HR, an expert recruitment analyst.
Analyze this job description:
---
{jd}
---
Candidate resume/profile context (optional):
---
{resume_context or 'Not provided'}
---
Return:
1. Role summary
2. Must-have requirements
3. Nice-to-have requirements
4. Technical / functional skills
5. Soft skills
6. Experience and education expectations
7. Important ATS keywords
8. Hidden or implied expectations
9. Candidate match assessment if profile was provided
10. Missing information / questions to ask the employer
11. Resume keywords and bullet themes the candidate should emphasize
Do not invent information.
"""
                    with st.spinner("Nexora is reading the job description..."): result=ai_text(prompt)
                    if result: st.session_state.jd_text=jd; st.session_state.jd_analysis=result; st.session_state.improved_jd_bytes=None; st.session_state.improved_jd_name=None; st.session_state.jd_fix_paid_demo=False
            if st.session_state.jd_analysis:
                with st.container(border=True):
                    st.markdown("### ✨ Make this JD recruitment-ready — one-time ₹149")
                    c1,c2=st.columns([1.0,2.0],vertical_alignment="center")
                    with c1:
                        if st.button("💳 Pay ₹149 (Temporary)",type="primary",use_container_width=True,key="jd_pay_demo"):
                            with st.spinner("Demo payment confirmed. Nexora is professionally rewriting the JD..."):
                                b,n=improve_job_description(st.session_state.jd_text,st.session_state.jd_analysis)
                            if b:
                                st.session_state.improved_jd_bytes=b; st.session_state.improved_jd_name=n; st.session_state.jd_fix_paid_demo=True
                    with c2:
                        st.write("AI improves structure, clarity, ATS readability and recruitment-ready wording without inventing facts.")
                        st.caption("Temporary demo checkout • No real payment is processed")
                st.divider(); st.subheader("📊 Job Description Assessment"); st.markdown(st.session_state.jd_analysis)
                if st.session_state.improved_jd_bytes:
                    st.success("Professional job description is ready.")
                    st.download_button("⬇️ Download Professional JD (Word)",data=st.session_state.improved_jd_bytes,file_name=st.session_state.improved_jd_name or "professional_job_description.docx",mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True,key="download_improved_jd")

        elif tool == "Offer Letter Analyzer":
            st.subheader("📑 Offer Letter Analyzer")
            st.write("Review an offer letter for salary structure, probation, notice period and clauses worth clarifying. Then create a professional decision-support pack.")
            offer=st.file_uploader("Upload offer letter",type=["pdf","png","jpg","jpeg"],key="offer_upload")
            if offer:
                valid,msg=validate_ai_file(offer)
                if not valid: st.error(msg)
                elif st.button("✦ Analyze Offer Letter",type="primary",use_container_width=True,key="analyze_offer_btn"):
                    prompt="""
You are Nexora HR, an employment-document review assistant for India.
Read the uploaded offer letter carefully and produce a plain-English decision-support report.
Return:
1. Executive summary
2. Employer / role / location / joining date
3. Fixed compensation
4. Variable compensation and performance-linked amounts
5. Deductions / benefits mentioned
6. Probation period
7. Notice period
8. Working hours / shifts / location conditions
9. Leave / benefits mentioned
10. Important restrictive or unusual clauses
11. Termination / separation clauses
12. Bond / service agreement / repayment clauses, if any
13. Confidentiality / IP / non-compete language, if any
14. Missing or unclear information
15. Questions the candidate should ask HR before accepting
16. Overall review: Favorable / Needs clarification / High attention, with reasons
Do not give legal advice or make unsupported legal conclusions. Do not invent information.
"""
                    with st.spinner("Nexora is reviewing the offer letter..."): result=ai_text(prompt,offer)
                    if result:
                        st.session_state.offer_analysis=result; st.session_state.offer_file_bytes=offer.getvalue(); st.session_state.offer_file_name=offer.name; st.session_state.offer_file_type=get_mime_type(offer); st.session_state.improved_offer_bytes=None; st.session_state.improved_offer_name=None; st.session_state.offer_fix_paid_demo=False
            if st.session_state.offer_analysis:
                with st.container(border=True):
                    st.markdown("### 📘 Create my Offer Review Pack — one-time ₹99")
                    c1,c2=st.columns([1.0,2.0],vertical_alignment="center")
                    with c1:
                        if st.button("💳 Pay ₹99 (Temporary)",type="primary",use_container_width=True,key="offer_pay_demo"):
                            upload=io.BytesIO(st.session_state.offer_file_bytes)
                            upload.name=st.session_state.offer_file_name or "offer.pdf"
                            upload.type=st.session_state.offer_file_type or "application/pdf"
                            with st.spinner("Demo payment confirmed. Nexora is preparing your review pack..."):
                                b,n=create_offer_review_pack(upload,st.session_state.offer_analysis)
                            if b:
                                st.session_state.improved_offer_bytes=b; st.session_state.improved_offer_name=n; st.session_state.offer_fix_paid_demo=True
                    with c2:
                        st.write("A professional decision-support pack with key terms, attention items, questions for HR and before-accepting checks.")
                        st.caption("Temporary demo checkout • No real payment is processed")
                st.divider(); st.subheader("📊 Offer Letter Assessment"); st.markdown(st.session_state.offer_analysis)
                if st.session_state.improved_offer_bytes:
                    st.success("Offer Review Pack is ready.")
                    st.download_button("⬇️ Download Offer Review Pack (Word)",data=st.session_state.improved_offer_bytes,file_name=st.session_state.improved_offer_name or "offer_letter_review_pack.docx",mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True,key="download_offer_pack")

        else:
            st.subheader("✉️ Cover Letter Generator")
            st.write("Create a tailored, professional cover letter from your resume and the job description. The generated letter stays grounded in your actual experience.")
            resume_text=st.text_area("Paste your resume text",height=230,placeholder="Paste your resume text here...",key="cover_resume_text")
            jd_text=st.text_area("Paste the job description",height=230,placeholder="Paste the job description here...",key="cover_jd_text")
            cover_role=st.text_input("Target role (optional)",key="cover_role")
            if st.button("✦ Generate Free Preview",type="primary",use_container_width=True,key="cover_preview_btn"):
                if len(resume_text.strip())<80 or len(jd_text.strip())<80: st.warning("Please provide both a fuller resume and job description for a useful tailored letter.")
                else:
                    prompt=f"""
Create a short preview of a tailored cover letter using only the candidate resume and job description below. Do not invent facts. Give 3 concise paragraphs and a suggested subject line.
Target role: {cover_role or 'Use the JD role'}
RESUME:\n{resume_text}\nJOB DESCRIPTION:\n{jd_text}
"""
                    with st.spinner("Nexora is preparing your cover letter preview..."): result=ai_text(prompt)
                    if result: st.session_state.cover_letter=result; st.session_state.cover_letter_bytes=None; st.session_state.cover_letter_paid_demo=False
            if st.session_state.cover_letter:
                with st.container(border=True):
                    st.markdown("### 🚀 Create the polished final letter — one-time ₹79")
                    c1,c2=st.columns([1.0,2.0],vertical_alignment="center")
                    with c1:
                        if st.button("💳 Pay ₹79 (Temporary)",type="primary",use_container_width=True,key="cover_pay_demo"):
                            with st.spinner("Demo payment confirmed. Nexora is creating the final cover letter..."):
                                b,n=generate_cover_letter(resume_text,jd_text,cover_role)
                            if b:
                                st.session_state.cover_letter_bytes=b; st.session_state.cover_letter_name=n; st.session_state.cover_letter_paid_demo=True
                    with c2:
                        st.write("AI tailors the final letter to the JD, improves wording and produces a polished editable Word document.")
                        st.caption("Temporary demo checkout • No real payment is processed")
                st.divider(); st.subheader("📝 Cover Letter Preview"); st.markdown(st.session_state.cover_letter)
                if st.session_state.cover_letter_bytes:
                    st.success("Your polished cover letter is ready.")
                    st.download_button("⬇️ Download Cover Letter (Word)",data=st.session_state.cover_letter_bytes,file_name=st.session_state.cover_letter_name or "tailored_cover_letter.docx",mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True,key="download_cover_letter")

    # ============================================================
    # HR TEMPLATES
    # ============================================================
    else:
        st.markdown("## 📚 HR Templates")
        st.caption("Professional, editable HR drafts with a clean corporate layout. Replace bracketed fields and review against your company policy before issue.")

        template = st.selectbox(
            "Choose a template",
            [
                "Offer Letter", "Appointment Letter", "Salary Increment Letter", "Promotion Letter",
                "Experience Certificate", "Relieving Letter", "Employee Warning Letter", "Exit Interview Form",
            ],
            key="hr_template_select",
        )

        templates = {
            "Offer Letter": {
                "subtitle":"Employment offer letter • editable draft",
                "metadata":[("Date","[DATE]"),("Reference","[REFERENCE NO.]"),("Candidate","[EMPLOYEE NAME]"),("Designation","[DESIGNATION]")],
                "sections":[
                    ("01  Offer of Employment",["Dear [EMPLOYEE NAME],","We are pleased to offer you employment with [COMPANY NAME] for the position of [DESIGNATION], subject to the terms set out in this letter and the applicable employment policies of the organization.","Proposed date of joining: [JOINING DATE]","Reporting location / work arrangement: [LOCATION / WORK ARRANGEMENT]"]),
                    ("02  Compensation & Benefits",["Annual CTC / gross compensation: [CTC]","The detailed compensation structure, benefits and applicable deductions will be communicated separately / attached as applicable."]),
                    ("03  Conditions of Employment",["This offer is subject to satisfactory completion of applicable background / reference checks, submission of required documents and compliance with company policies.","The final employment relationship will be governed by the appointment / employment documentation issued by the company and applicable law."]),
                    ("04  Acceptance",["Please sign and return a copy of this letter by [ACCEPTANCE DATE] to confirm your acceptance of the offer.","We look forward to welcoming you to [COMPANY NAME]."]),
                ]},
            "Appointment Letter": {
                "subtitle":"Formal appointment letter • editable draft",
                "metadata":[("Date","[DATE]"),("Employee ID","[EMPLOYEE ID]"),("Employee","[EMPLOYEE NAME]"),("Designation","[DESIGNATION]")],
                "sections":[
                    ("01  Appointment",["Dear [EMPLOYEE NAME],","You are appointed as [DESIGNATION] in [DEPARTMENT] with effect from [DATE]. We are pleased to have you join [COMPANY NAME]."]),
                    ("02  Place & Nature of Work",["Primary work location: [LOCATION]","Reporting manager: [MANAGER]","The role, responsibilities and reasonable work arrangements may be updated in line with business requirements and company policy."]),
                    ("03  Compensation",["Your annual CTC / gross compensation will be [CTC]. The detailed salary structure and eligible benefits are as communicated separately."]),
                    ("04  Probation",["Probation period: [PERIOD]. Confirmation will be subject to performance, conduct, business requirements and the applicable company policy."]),
                    ("05  Notice & Separation",["Notice period: [DAYS] days, subject to the employment contract, company policy and applicable law."]),
                    ("06  Acknowledgement",["Please sign below to acknowledge receipt and acceptance of the appointment terms."]),
                ]},
            "Salary Increment Letter": {
                "subtitle":"Compensation revision communication • editable draft",
                "metadata":[("Date","[DATE]"),("Employee ID","[EMPLOYEE ID]"),("Employee","[EMPLOYEE NAME]"),("Designation","[DESIGNATION]")],
                "sections":[
                    ("01  Compensation Revision",["Dear [EMPLOYEE NAME],","We are pleased to inform you that, in recognition of your contribution and performance, your compensation has been revised with effect from [EFFECTIVE DATE]."]),
                    ("02  Compensation Summary",["Previous annual CTC / gross: [OLD CTC]","Revised annual CTC / gross: [NEW CTC]","Increment: [PERCENTAGE]%"]),
                    ("03  Closing",["We appreciate your contribution to [COMPANY NAME] and look forward to your continued growth and success."]),
                ]},
            "Promotion Letter": {
                "subtitle":"Promotion communication • editable draft",
                "metadata":[("Date","[DATE]"),("Employee ID","[EMPLOYEE ID]"),("Employee","[EMPLOYEE NAME]"),("Effective Date","[EFFECTIVE DATE]")],
                "sections":[
                    ("01  Promotion",["Dear [EMPLOYEE NAME],","We are pleased to inform you that you have been promoted to [NEW DESIGNATION], effective [EFFECTIVE DATE], in recognition of your contribution and performance."]),
                    ("02  Role Details",["Department / Function: [DEPARTMENT]","Reporting manager: [MANAGER]","Primary responsibilities: [BRIEF ROLE SUMMARY]"]),
                    ("03  Compensation",["Revised annual CTC / gross: [CTC]","The detailed compensation structure and applicable benefits will be communicated separately."]),
                    ("04  Congratulations",["Congratulations on this achievement. We wish you continued success in your expanded role."]),
                ]},
            "Experience Certificate": {
                "subtitle":"Employment experience certificate • editable draft",
                "metadata":[("Date","[DATE]"),("Reference","[REFERENCE NO.]"),("Employee","[EMPLOYEE NAME]"),("Employee ID","[EMPLOYEE ID]")],
                "sections":[
                    ("01  To Whom It May Concern",["This is to certify that [EMPLOYEE NAME] was employed with [COMPANY NAME] as [DESIGNATION] from [START DATE] to [END DATE]."]),
                    ("02  Employment Details",["Department / Function: [DEPARTMENT / FUNCTION]","Location: [LOCATION]","The above information is issued based on the employment records available with the organization."]),
                    ("03  Closing",["We wish [EMPLOYEE NAME] all the best in future professional endeavors."]),
                ]},
            "Relieving Letter": {
                "subtitle":"Employee separation confirmation • editable draft",
                "metadata":[("Date","[DATE]"),("Reference","[REFERENCE NO.]"),("Employee","[EMPLOYEE NAME]"),("Employee ID","[EMPLOYEE ID]")],
                "sections":[
                    ("01  Relieving Confirmation",["Dear [EMPLOYEE NAME],","This is to confirm that you have been relieved from your services with [COMPANY NAME] effective [LAST WORKING DATE], following completion of the applicable separation formalities."]),
                    ("02  Handover & Clearance",["The required handover, asset return and clearance activities have been completed / are recorded as applicable in the organization's records."]),
                    ("03  Closing",["We thank you for your services and contribution during your association with [COMPANY NAME]. We wish you success in your future endeavors."]),
                ]},
            "Employee Warning Letter": {
                "subtitle":"Formal workplace warning • editable draft",
                "metadata":[("Date","[DATE]"),("Reference","[REFERENCE NO.]"),("Employee","[EMPLOYEE NAME]"),("Employee ID","[EMPLOYEE ID]")],
                "sections":[
                    ("01  Subject",["Subject: Formal Warning — [SUBJECT]"]),
                    ("02  Concern / Incident",["This letter is being issued regarding [FACTUAL DESCRIPTION OF ISSUE], observed / reported on [DATE(S)]. The concern has been discussed / documented as applicable."]),
                    ("03  Expected Improvement",["You are required to [CLEAR EXPECTATION] with immediate effect and to maintain the expected standard consistently.","Any supporting documents, examples or relevant policy references may be attached to this letter."]),
                    ("04  Review & Support",["Progress will be reviewed on [REVIEW DATE / PERIOD]. Please contact [MANAGER / HR] if you require clarification regarding the expectations set out above."]),
                    ("05  Important Review",["This draft should be adapted to the organization's disciplinary process, documented facts, applicable policy and applicable law before issue."]),
                ]},
            "Exit Interview Form": {
                "subtitle":"Structured employee exit feedback form",
                "metadata":[("Date","[INTERVIEW DATE]"),("Employee","[EMPLOYEE NAME]"),("Department","[DEPARTMENT]"),("Last Working Day","[DATE]")],
                "sections":[
                    ("01  Interview Details",["Interviewer: [INTERVIEWER NAME]","Designation: [DESIGNATION]","Tenure: [TENURE]"]),
                    ("02  Key Discussion Points",["Primary reason for leaving: [REASON]","Has another offer been accepted? [YES / NO / PREFER NOT TO SAY]","What did you value most about your experience? [RESPONSE]","What should the organization improve? [RESPONSE]","What could have influenced you to stay? [RESPONSE]"]),
                    ("03  Ratings",["Use the structured rating table below and add comments where useful."]),
                    ("04  Additional Notes",["HR / interviewer notes: [NOTES]"]),
                ]},
        }

        item = templates[template]
        c1, c2 = st.columns([1.45, 1], gap="small", vertical_alignment="top")
        with c1:
            with st.container(border=True):
                st.subheader(f"📝 {template}")
                st.caption(item["subtitle"])
                st.markdown("**Professional document preview**")
                _metadata = item["metadata"]
                _sections = item["sections"]
                for label, value in _metadata:
                    st.write(f"**{label}:** {value}")
                st.divider()
                for heading, lines in _sections:
                    st.markdown(f"**{heading}**")
                    for line in lines:
                        st.write(line)
                if template == "Exit Interview Form":
                    st.markdown("**Rating areas:** Overall experience · Manager support · Work environment · Growth · Compensation · Reason for leaving · Would recommend")
        with c2:
            with st.container(border=True):
                st.subheader("⬇️ Professional Word draft")
                st.write("Clean corporate formatting, consistent typography, structured sections and signature area — ready to customize.")
                content = docx_bytes(template, item["subtitle"], item["sections"], metadata=item["metadata"], form=(template == "Exit Interview Form"))
                filename = re.sub(r"[^a-z0-9]+", "_", template.lower()).strip("_") + "_template.docx"
                st.download_button(
                    "⬇️ Download Editable Word File",
                    data=content,
                    file_name=filename,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary",
                    use_container_width=True,
                )
                st.success("✓ No Nexora promotional branding is inserted into the downloaded document.")
                st.info("Before issuing: replace all [BRACKETED] fields, verify names/dates/amounts, and review against company policy and applicable law.")

    # ============================================================
    st.divider()
    with st.container(border=True):
        st.subheader("🚀 Nexora HR roadmap")
        r1, r2, r3, r4 = st.columns(4, gap="small")
        for col, title, text in [
            (r1, "NOW", "Free salary calculators + HR templates + AI career analysis"),
            (r2, "NEXT", "SEO landing pages and shareable result pages"),
            (r3, "PRO", "Detailed reports, saved history and higher AI limits"),
            (r4, "SCALE", "HR teams, recruiter tools and business subscriptions"),
        ]:
            with col:
                st.markdown(f"**{title}**")
                st.caption(text)

    st.caption("✦ NEXORA HR  ·  Practical tools for work and career")

# ============================================================
# MAIN EXPERIENCE
# ============================================================
if st.session_state.main_section == "HR":
    # A clear return path prevents HR from feeling like a duplicate home page.
    back_col, title_col = st.columns([1.0, 5.5], gap="small", vertical_alignment="center")
    with back_col:
        if st.button("← Home", use_container_width=True, key="hr_back_home"):
            st.session_state.main_section = "Home"
            st.session_state.home_choice = None
            st.rerun()
    with title_col:
        st.caption("NEXORA WORKSPACE")
    render_hr_hub()

elif st.session_state.main_section == "Documents" and not st.session_state.document_ready:
    # ------------------------------------------------------------
    # DOCUMENT AI LANDING / WORKSPACE
    # ------------------------------------------------------------
    back_col, title_col = st.columns([1.0, 5.5], gap="small", vertical_alignment="center")
    with back_col:
        if st.button("← Home", use_container_width=True, key="docs_back_home"):
            st.session_state.main_section = "Home"
            st.session_state.home_choice = None
            st.rerun()
    with title_col:
        st.caption("NEXORA DOCUMENT AI INTELLIGENCE")

    st.markdown("# Your Documents — **Smarter With Nexora**")
    st.write("Choose the document capability you need, then upload your file and let Nexora do the work.")
    st.caption("🔒 Documents are processed by Nexora using Google Gemini AI. Avoid uploading information you are not authorized to share.")

    # Clickable document capabilities — no radio pills.
    st.markdown("### Choose What You Want To Do")
    tool_cards = [
        ("🧠", "AI Summary", "Understand the key points quickly.", "Summary"),
        ("💬", "Document Chat", "Ask questions directly about your file.", "Chat"),
        ("📊", "Data Extraction", "Turn document content into structured data.", "Extract Data"),
        ("🔍", "Deep Analysis", "Find risks, gaps and important details.", "Analyze"),
    ]
    for row in range(0, len(tool_cards), 4):
        cols = st.columns(4, gap="small")
        for col, (icon, title, text, value) in zip(cols, tool_cards[row:row+4]):
            with col:
                with st.container(border=True):
                    if st.button(f"{icon}  {title}", use_container_width=True, key=f"doc_choice_{value}"):
                        st.session_state.document_tool_choice = value
                        st.rerun()
                    st.caption(text)

    selected_tool = st.session_state.document_tool_choice
    st.markdown(f"### Upload For: {selected_tool}")

    upload_col, info_col = st.columns([1.55, 1], gap="small", vertical_alignment="top")
    with upload_col:
        with st.container(border=True):
            st.subheader("☁️ Start With A Document")
            st.caption("PDF, PNG, JPG or JPEG · Up to 50 MB")
            uploaded_file = st.file_uploader(
                "Choose a document",
                type=["pdf", "png", "jpg", "jpeg"],
                label_visibility="collapsed",
                max_upload_size=50,
                key="home_document_upload",
            )
            if uploaded_file:
                valid, message = validate_file(uploaded_file)
                if not valid:
                    st.error(message)
                else:
                    st.success(f"✓ {uploaded_file.name} is ready")
                    if st.button("✦  Analyze Document", type="primary", use_container_width=True, key="home_analyze"):
                        if process_document(uploaded_file):
                            # Open the result the user selected from the home page.
                            choice = st.session_state.document_tool_choice
                            st.session_state.active_view = "Summary" if choice == "Summary" else choice
                            st.rerun()
            else:
                st.caption("Upload once. Nexora will keep the document visible while you work.")

    with info_col:
        with st.container(border=True):
            st.subheader("✦ Included In Document AI")
            for icon, title, text in [
                ("🧠", "AI Summary", "Executive summary and key takeaways."),
                ("💬", "Document Chat", "Natural-language Q&A about the file."),
                ("📊", "Data Extraction", "Names, dates, amounts and line items."),
                ("🔍", "Deep Analysis", "Risks, missing information and inconsistencies."),
                ("📄", "Document View", "Keep the source visible while working."),
            ]:
                st.markdown(f"**{icon} {title}**")
                st.caption(text)

elif st.session_state.main_section == "Documents" and st.session_state.document_ready:
    header, home_col, new_doc = st.columns([4.7, 1.0, 1.15], gap="small", vertical_alignment="center")
    with header:
        st.subheader(f"📄 {st.session_state.document_name}")
        st.caption("● Document analyzed and ready")
    with home_col:
        if st.button("← Home", use_container_width=True, key="ready_home"):
            st.session_state.main_section = "Home"
            st.session_state.home_choice = None
            st.rerun()
    with new_doc:
        if st.button("＋ New", use_container_width=True, key="new_document"):
            reset_workspace()
            st.session_state.main_section = "Documents"
            st.rerun()

    # Clickable tool navigation instead of round radio selections.
    st.markdown("### Document Workspace")
    tool_nav = [
        ("✨ Summary", "Summary"),
        ("💬 Chat", "Chat"),
        ("📊 Extract Data", "Extract Data"),
        ("🔍 Analyze", "Analyze"),
    ]
    nav_cols = st.columns(4, gap="small")
    for col, (label, value) in zip(nav_cols, tool_nav):
        with col:
            if st.button(label, use_container_width=True, key=f"ready_tool_{value}"):
                st.session_state.document_tool_choice = value
                st.session_state.active_view = value
                st.rerun()

    st.divider()

    document_column, chat_column = st.columns([1.65, 1], gap="small", vertical_alignment="top")
    with document_column:
        with st.container(border=True, height=640):
            st.subheader("📄 Document")
            st.caption(st.session_state.document_name)
            st.divider()
            if st.session_state.document_type == "application/pdf":
                try:
                    st.pdf(st.session_state.document_bytes, height=340)
                except Exception:
                    st.info("PDF preview is unavailable. The document is still available to Nexora.")
            elif st.session_state.document_type and st.session_state.document_type.startswith("image/"):
                st.image(st.session_state.document_bytes, use_container_width=True)

            st.divider()

            if st.session_state.active_view == "Summary":
                st.subheader("✨ AI Summary")
                if st.session_state.summary:
                    st.markdown(st.session_state.summary)

            elif st.session_state.active_view == "Extract Data":
                st.subheader("📊 Extracted Data")
                if st.session_state.extracted_information is None and st.session_state.extracted_line_items is None:
                    st.info("Extract structured information from this document.")
                    if st.button("📊  Extract Data", type="primary", use_container_width=True, key="extract_data_btn"):
                        extract_data()
                else:
                    info = st.session_state.extracted_information or []
                    items = st.session_state.extracted_line_items or []
                    if info:
                        st.write("**Document Information**")
                        st.dataframe(pd.DataFrame(info), use_container_width=True, hide_index=True)
                    if items:
                        st.write("**Line Items**")
                        st.dataframe(pd.DataFrame(items), use_container_width=True, hide_index=True)
                    if info or items:
                        st.download_button(
                            "⬇️  Download Excel",
                            data=create_excel(),
                            file_name="nexora_extracted_data.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True,
                        )
                    else:
                        st.warning("No structured information was found.")

            elif st.session_state.active_view == "Analyze":
                st.subheader("🔍 Deep Analysis")
                if not st.session_state.analysis_result:
                    st.info("Run deeper analysis to identify risks, missing information, inconsistencies and next steps.")
                    if st.button("🔍  Run Deep Analysis", type="primary", use_container_width=True, key="deep_analysis_btn"):
                        run_deep_analysis()
                else:
                    st.markdown(st.session_state.analysis_result)

    with chat_column:
        with st.container(border=True):
            st.subheader("✦ Nexora AI")
            st.caption("Ask questions and work with your document in real time.")

        with st.container(height=575, border=True):
            if not st.session_state.messages:
                st.info("Ask Nexora anything about this document.")
                st.write("**Try Asking:**")
                suggestions = [
                    "What is this document about?",
                    "What are the most important points?",
                    "Are there any risks I should know about?",
                    "What are the important dates and amounts?",
                    "Summarize this in simple language.",
                ]
                for idx, suggestion in enumerate(suggestions):
                    if st.button(suggestion, use_container_width=True, key=f"chat_suggestion_{idx}"):
                        st.session_state.messages.append({"role": "user", "content": suggestion})
                        with st.chat_message("assistant"):
                            placeholder = st.empty()
                            answer = ask_document(suggestion, placeholder)
                        if answer:
                            st.session_state.messages.append({"role": "assistant", "content": answer})
                        st.rerun()
            else:
                for message in st.session_state.messages:
                    with st.chat_message(message["role"]):
                        st.markdown(message["content"])

        question = st.chat_input("Ask anything about this document...", key="nexora_chat_input")
        if question:
            st.session_state.messages.append({"role": "user", "content": question})
            with st.chat_message("user"):
                st.markdown(question)
            with st.chat_message("assistant"):
                placeholder = st.empty()
                answer = ask_document(question, placeholder)
            if answer:
                st.session_state.messages.append({"role": "assistant", "content": answer})
            st.rerun()

else:
    # ------------------------------------------------------------
    # HOME — USER CHOOSES THE PRODUCT AREA
    # ------------------------------------------------------------
    st.markdown("# Welcome To Nexora ✦")
    st.write("Choose the workspace you want. Nexora will take you directly to the tools for that job.")

    st.markdown("## Choose Your Nexora Workspace")
    choice_left, choice_right = st.columns(2, gap="small", vertical_alignment="top")

    with choice_left:
        with st.container(border=True):
            st.markdown("# 🧠")
            st.markdown("## Document AI Intelligence")
            st.write("Understand, chat with, extract data from and deeply analyze your documents.")
            if st.button("Open Document AI →", type="primary", use_container_width=True, key="home_document_choice"):
                st.session_state.home_choice = "Documents"
                st.session_state.main_section = "Documents"
                st.session_state.document_tool_choice = "Summary"
                st.rerun()
            st.caption("AI Summary · Document Chat · Data Extraction · Deep Analysis · Document View")

    with choice_right:
        with st.container(border=True):
            st.markdown("# 💼")
            st.markdown("## HR & Career Intelligence")
            st.write("Work with salary tools, career analysis, recruitment documents and professional HR templates.")
            if st.button("Open HR & Career →", type="primary", use_container_width=True, key="home_hr_choice"):
                st.session_state.home_choice = "HR"
                st.session_state.main_section = "HR"
                st.session_state.hr_section = "Salary & HR Calculators"
                st.rerun()
            st.caption("Salary Tools · Resume · Job Description · Offer Letter · Cover Letter · HR Templates")

    st.divider()
    st.markdown("## Explore All Nexora Capabilities")
    all_tools = [
        ("🧠", "AI Summary", "Turn long documents into clear key points."),
        ("💬", "Document Chat", "Ask natural-language questions about a file."),
        ("📊", "Data Extraction", "Convert important content into structured tables."),
        ("🔍", "Deep Analysis", "Surface risks, gaps, dates and important details."),
        ("💰", "Salary Calculators", "Estimate take-home, increments, gratuity and notice pay."),
        ("📄", "Resume Analyzer", "Identify resume strengths, gaps and ATS issues."),
        ("🎯", "JD Analyzer", "Review job descriptions and improve clarity."),
        ("📝", "Cover Letter", "Create a polished role-specific application letter."),
        ("📋", "Offer Review", "Review an offer letter and create a structured pack."),
        ("🗂️", "HR Templates", "Download professional editable HR documents."),
    ]
    for start_row in range(0, len(all_tools), 5):
        cols = st.columns(5, gap="small")
        for col, (icon, title, text) in zip(cols, all_tools[start_row:start_row+5]):
            with col:
                with st.container(border=True):
                    st.markdown(f"### {icon} {title}")
                    st.caption(text)

    st.caption("✦ One home. Two focused workspaces. No duplicated navigation.")

st.caption("✦ NEXORA  ·  AI Document Intelligence • HR & Career Intelligence")
